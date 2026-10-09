"""Verify the UI-selected Holt–Winters run without changing source files."""
import json
import math
import sys
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
base = json.loads((root/'runs/5d6320f4d264/result.json').read_text())
identifier = sys.argv[1] if len(sys.argv) > 1 else '2c94a1cc0dff'
if len(identifier) != 12 or any(c not in '0123456789abcdef' for c in identifier):
    raise SystemExit('Use a saved forecast run identifier.')
selected = json.loads((root/'runs'/identifier/'result.json').read_text())
assert selected['method_selection'] in {'model:Holt-Winters seasonal', 'seasonal'}
assert selected['metrics']['evaluation_signature'] == base['metrics']['evaluation_signature']
assert selected['input_manifest']['sources'] == base['input_manifest']['sources']
for item in selected['items']:
    weights = selected['series_ensemble_weights'][item]
    if selected['method_selection'].startswith('model:'):
        assert weights == {'Holt-Winters seasonal': 1.0}
    else:
        assert set(weights).issubset({'Holt-Winters seasonal', 'Seasonal naive', 'AutoETS', 'Theta', 'AutoARIMA', 'MSTL weekly + yearly'})
    assert math.isclose(math.fsum(weights.values()), 1.)
    assert selected['series'][item]['history'] == base['series'][item]['history']
    for index, forecast in enumerate(selected['series'][item]['forecast']):
        expected = math.fsum(weight * selected['series'][item]['methods'][name][index]['mean'] for name, weight in weights.items())
        assert math.isclose(forecast['mean'], expected)
for original in base['leaderboard']:
    counterpart = next(r for r in selected['leaderboard'] if r['model'] == original['model'])
    assert original['wape_pct'] == counterpart['wape_pct']
book = openpyxl.load_workbook(root/'runs'/identifier/'forecast_package.xlsx', data_only=True)
assert dict(book['Engine'].iter_rows(min_row=2, values_only=True)) == selected['engine']
assert {name: json.loads(settings) for name, settings in book['Method settings'].iter_rows(min_row=2, values_only=True)} == selected['method_settings']
assert book['Forecast'].max_row - 1 == 144
book.close()
print(json.dumps({'run_id': selected['run_id'], 'method': selected['best_model'], 'items': len(selected['items']),
    'forecast_rows': len(selected['forecast_rows']), 'later_period_wape_pct': selected['metrics']['wape_pct'],
    'forecast_total': math.fsum(r['mean'] for r in selected['series']['__all__']['forecast']),
    'same_history_source_hashes_and_test_periods': True, 'method_settings_export_matches': True}))
