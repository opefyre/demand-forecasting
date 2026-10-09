"""Small local sample acceptance. --start creates one labelled synthetic baseline.

--verify RUN_ID checks a completed no-factor comparison, totals and exports.
Never changes original client files, order snapshots or earlier forecasts.
"""
import argparse
from copy import deepcopy
from io import BytesIO
import json
import math
import sys
from pathlib import Path

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import httpx
import openpyxl

BASE='http://127.0.0.1:8010'
def api(path,payload=None):
    response=httpx.get(BASE+path,timeout=30) if payload is None else httpx.post(BASE+path,json=payload,timeout=30)
    response.raise_for_status();return response.json()

def start():
    source=next(d for d in api('/api/datasets')['datasets'] if d['name']=='First-upload review · sample' and d['classification']=='synthetic_sample')
    settings=deepcopy(source['settings'])
    settings.update(drivers=['usd_irr_synthetic','energy_curtailment_hours_synthetic'],
        method_selection='model:Ridge + drivers',horizon=6,profile='fast',
        factor_definitions={
            'usd_irr_synthetic':{'unit':'IRR per USD','geography':'Iran · synthetic market','source':'Bundled synthetic sample, not live FX'},
            'energy_curtailment_hours_synthetic':{'unit':'hours','geography':'Synthetic site, not client observations','source':'Bundled synthetic sample'}})
    saved=api('/api/datasets',{'name':'Factor comparison · synthetic baseline','sources':{k:v for k,v in source['sources'].items() if k in {'history','future'}},
        'settings':settings,'classification':'synthetic_sample','accept_warnings':True,'request_id':'factor-cutoff-demo-baseline-20260924-1'})
    print(json.dumps(api('/api/jobs',{'dataset_id':saved['id'],'request_id':'factor-cutoff-demo-job-20260924-1'})))

def verify(key):
    comparison=api('/api/runs/'+key);base=api('/api/runs/'+comparison['base_run_id'])
    assert comparison['scenario']['type']=='factor_comparison'
    assert base['engine']==comparison['engine']
    assert base['metrics']['evaluation_signature']==comparison['metrics']['evaluation_signature']
    assert base['metrics']['factor_test_policy']=='last_training_value'
    assert comparison['metrics']['factor_test_policy']=='no_extra_factors'
    assert base['input_manifest']['sources']==comparison['input_manifest']['sources']
    assert base['input_manifest']['settings']['drivers'] and not comparison['input_manifest']['settings']['drivers']
    assert base['source_classification']==comparison['source_classification']=='synthetic_sample'
    for item in base['items']:
        assert base['series'][item]['history']==comparison['series'][item]['history']
    # Independently reconcile aggregate forecasts against customer–SKU rows.
    for run in (base,comparison):
        for total in run['series']['__all__']['forecast']:
            expected=math.fsum(point['mean'] for item in run['items'] for point in run['series'][item]['forecast'] if point['timestamp']==total['timestamp'])
            assert math.isclose(total['mean'],expected,rel_tol=1e-9,abs_tol=1e-6)
    response=httpx.get(BASE+'/api/export/'+key+'/xlsx');response.raise_for_status()
    book=openpyxl.load_workbook(BytesIO(response.content),data_only=True)
    assert book['Factor comparison'].max_row==3
    settings=dict(book['Run Settings'].iter_rows(min_row=2,values_only=True))
    assert settings['factor_test_policy']=='no_extra_factors'
    print(json.dumps({'baseline_run':base['run_id'],'comparison_run':key,
        'with_factors_error_pct':base['metrics']['wape_pct'],'without_factors_error_pct':comparison['metrics']['wape_pct'],
        'baseline_total':sum(r['mean'] for r in base['series']['__all__']['forecast']),
        'comparison_total':sum(r['mean'] for r in comparison['series']['__all__']['forecast']),
        'series':len(base['items']),'export_verified':True,'classification':'synthetic_sample'}))

if __name__=='__main__':
    parser=argparse.ArgumentParser();group=parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--start',action='store_true');group.add_argument('--verify')
    args=parser.parse_args();start() if args.start else verify(args.verify)
