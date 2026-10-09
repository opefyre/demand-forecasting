"""Independent saved-run arithmetic; does not import the range implementation."""
import json
import math
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parents[1]
identifier = sys.argv[1] if len(sys.argv) > 1 else '25ad55031c82'
if len(identifier) != 12 or any(c not in '0123456789abcdef' for c in identifier):
    raise SystemExit('Use a saved forecast run identifier.')
run = json.loads((root/'runs'/identifier/'result.json').read_text())
model, check = run['range_model'], run['metrics']['range_check']
assert run['source_classification'] == 'synthetic_sample'
selection, fitting, later = (set(run['metrics'][key]) for key in ('selection_periods', 'range_fitting_periods', 'confirmation_periods'))
assert not selection & fitting and not fitting & later and not selection & later
assert run['metrics']['rolling_folds'] == 5
assert len(selection) == len(fitting) == 12 and len(later) == 6
errors = defaultdict(list)
joint = defaultdict(list)
for row in run['range_fitting_rows']:
    assert row['timestamp'] in fitting
    errors[row['item_id']].append(abs(row['actual'] - row['predicted']))
    joint[row['timestamp']].append(row)
for date, group in joint.items():
    assert {r['item_id'] for r in group} == set(run['items'])
    errors['__portfolio__'].append(abs(math.fsum(r['actual'] - r['predicted'] for r in group)))
widths = {item: sorted(values)[math.ceil((len(values)+1)*.8)-1] for item, values in errors.items()}
for item, parameters in model['parameters'].items():
    for parameter in parameters.values():
        assert parameter['scope'] == 'pooled_horizons'
        assert math.isclose(parameter['half_width'], widths[item], rel_tol=1e-10, abs_tol=1e-10)
checked_items, checked_totals = [], []
for row in check['rows']:
    assert row['timestamp'] in later
    half = widths[row['item_id']]
    assert math.isclose(row['lower'], max(0, row['predicted'] - half), abs_tol=1e-9)
    assert math.isclose(row['upper'], row['predicted'] + half, abs_tol=1e-9)
    inside = row['lower'] <= row['actual'] <= row['upper']
    assert inside == row['covered']
    (checked_totals if row['item_id'] == '__portfolio__' else checked_items).append(inside)
for name, values in [('items', checked_items), ('portfolio', checked_totals)]:
    assert check[name]['checked'] == len(values)
    assert math.isclose(check[name]['coverage_pct'], 100*sum(values)/len(values))
for item, series in run['series'].items():
    half = widths['__portfolio__' if item == '__all__' else item]
    for row in series['forecast']:
        assert math.isclose(row['p10'], max(0, row['mean']-half), abs_tol=1e-9)
        assert math.isclose(row['p90'], row['mean']+half, abs_tol=1e-9)
book = openpyxl.load_workbook(root/'runs'/identifier/'forecast_package.xlsx', data_only=True)
assert book['Range fitting'].max_row - 1 == len(run['range_fitting_rows'])
sheet = book['Range check']; headers = [c.value for c in sheet[1]]
exported = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
assert len(exported) == len(check['rows'])
for actual, expected in zip(exported, check['rows']):
    for key, value in actual.items():
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            assert math.isclose(value, expected[key], abs_tol=1e-9)
        else:
            assert value == expected[key]
book.close()
print(json.dumps({'run_id': identifier, 'revision': run['engine']['revision'],
    'range_fitting_rows': len(run['range_fitting_rows']), 'item_coverage': check['items'],
    'portfolio_coverage': check['portfolio'], 'joint_width': widths['__portfolio__'],
    'saved_bounds_and_export_reconciled': True}))
