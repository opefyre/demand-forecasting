"""Read-only reconciliation of the browser-created synthetic plan."""
import json
import hashlib
import math
from io import BytesIO
from pathlib import Path
import sys

import httpx
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.datasets import DatasetStore
from app.operations import read_operations_workbook

root = Path(__file__).resolve().parents[1]
identifier = 'e2379cd021'
with httpx.Client(base_url='http://127.0.0.1:8010', timeout=30) as client:
    plan = next(p for p in client.get('/api/plans').json()['plans'] if p['id'] == identifier)
    assert plan['status'] == 'published'
    assert plan['settings']['source_classification'] == 'synthetic_sample'
    review = client.get(f'/api/plans/{identifier}/quantities').json()
    supply = client.get(f'/api/plans/{identifier}/supply').json()
    response = client.get(f'/api/plans/{identifier}/export?include_supply=true')
    response.raise_for_status()
run = json.loads((root/'runs'/plan['run_id']/'result.json').read_text())
assert plan['run_id'] == '431fed798551'
assert len(review['rows']) == 72
changed = [r for r in review['rows'] if r['adjustment'] != 0]
assert len(changed) == 1
assert changed[0]['item_id'] == 'COA-ART-135 · Domestic' and changed[0]['period'] == '2026-09-01'
assert changed[0]['plan_quantity'] == 125
assert run['series'][changed[0]['item_id']]['forecast'][0]['mean'] == changed[0]['forecast_quantity']
assert len(plan['overrides']) == 2 and plan['overrides'][0]['reverted_at']
book = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
for sheet_name, expected in [('Plan quantities', review['rows']), ('Materials', supply['materials'])]:
    sheet = book[sheet_name]
    headers = [c.value for c in sheet[1]]
    actual = [dict(zip(headers, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
    assert len(actual) == len(expected)
    for a, e in zip(actual, expected):
        for k, value in a.items():
            if isinstance(value, (int, float)):
                assert math.isclose(value, e[k], abs_tol=1e-9)
            else:
                assert (value or '') == (e[k] or '')
book.close()
source = next(s for s in run['input_manifest']['sources'] if s['role'] == 'operations')
stored, content = DatasetStore(root/'data/datasets').source(source['id'])
assert stored['sha256'] == source['sha256']
tables = read_operations_workbook(stored['name'], content, run['input_manifest']['settings']['operations_mapping'])
recipe = {r.sku: float(r.quantity_per_tonne)*(1 + float(r.scrap_pct)/100)
          for r in tables['bom'].itertuples() if r.material_id == 'PULP-HW'}
required = math.fsum(r['plan_quantity']*recipe.get(r['sku'],0) for r in review['rows'] if r['period'] == '2026-09-01')
material = next(r for r in supply['materials'] if r['material_id'] == 'PULP-HW' and r['period'] == '2026-09-01')
assert material['gross_requirement'] == round(required,3)
assert supply['quantity_basis']['id'] == identifier
print(json.dumps({'plan_id': identifier, 'status': plan['status'], 'review_and_export_rows': len(review['rows']),
    'baseline_quantity': changed[0]['forecast_quantity'], 'plan_quantity': 125,
    'adjustment': changed[0]['adjustment'], 'pulp_requirement_independent': required,
    'pulp_requirement_supply': material['gross_requirement'], 'source_unchanged': True}))

assert hashlib.sha256(json.dumps(plan, sort_keys=True).encode()).hexdigest() == 'b72ea8f06e59524a67aeb2b1ea2fab8f01378cac586c9ee3e28cdf06bca83f85'
with httpx.Client(base_url='http://127.0.0.1:8010', timeout=30) as client:
    revision = client.get('/api/plans/51f34a7686/quantities').json()
    response = client.get('/api/plans/51f34a7686/export')
    response.raise_for_status()
changes = [r for r in revision['rows'] if r['revision_change']]
assert len(changes) == 1 and changes[0]['previous_plan_quantity'] == 125
assert changes[0]['plan_quantity'] == 135 and changes[0]['revision_change'] == 10
book = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
values = list(book['Plan quantities'].values)
assert len(values)-1 == len(revision['rows']) == 72
for exported, expected in zip(values[1:], revision['rows']):
    for key, value in zip(values[0], exported):
        if isinstance(value, (int, float)):
            assert math.isclose(value, expected[key], abs_tol=1e-9)
        else:
            assert (value or '') == (expected[key] or '')
assert dict(book['Plan version'].iter_rows(min_row=2, values_only=True))['parent_plan_id'] == identifier
book.close()
print(json.dumps({'revision_id': '51f34a7686', 'parent_record_unchanged': True, 'previous_quantity': 125,
                  'revised_quantity': 135, 'revision_change': 10, 'export_rows_reconciled': 72}))
