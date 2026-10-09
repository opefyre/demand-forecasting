"""Read-only workbook audit. Never evaluates formulas or changes the input files."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re

import openpyxl


def audit(path: Path) -> dict:
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
              "formula_calculation": "Saved Excel values only; formulas not recalculated", "sheets": []}
    try:
        for sheet in formulas:
            report = {"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                      "formulas": 0, "uncached_formulas": 0, "external_formulas": 0,
                      "error_counts": Counter(), "error_examples": [], "formula_examples": {},
                      "headers": [], "date_headers": []}
            functions = Counter()
            for frow, vrow in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                for cell, value_cell in zip(frow, vrow):
                    if cell.data_type == "f":
                        report["formulas"] += 1
                        report["uncached_formulas"] += value_cell.value is None
                        formula = str(cell.value)
                        report["external_formulas"] += bool(re.search(r"\[\d+\]", formula))
                        names = re.findall(r"([A-Z][A-Z0-9_.]*)\s*\(", formula.upper()) or ["arithmetic/reference"]
                        functions.update(names)
                        for name in names:
                            report["formula_examples"].setdefault(name, {"cell": cell.coordinate,
                                "formula": formula[:600], "saved_value": value_cell.value})
                    if value_cell.data_type == "e":
                        report["error_counts"][str(value_cell.value)] += 1
                        if len(report["error_examples"]) < 8:
                            report["error_examples"].append({"cell": cell.coordinate, "error": value_cell.value})
                    if getattr(cell, 'row', 100) <= 10:
                        if isinstance(value_cell.value, (datetime, date)):
                            report["date_headers"].append({"cell": cell.coordinate, "date": value_cell.value.isoformat()})
                        elif isinstance(value_cell.value, str) and cell.data_type != "f" and value_cell.data_type != "e":
                            report["headers"].append({"cell": cell.coordinate, "value": value_cell.value[:100]})
            report["functions"] = dict(functions.most_common())
            result["sheets"].append(report)
    finally:
        formulas.close()
        cached.close()
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    reports = [audit(path) for path in args.files]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(reports, ensure_ascii=False, indent=2, default=str))
    for report in reports:
        print(report['file'], report['sha256'])
        for sheet in report['sheets']:
            print(sheet['name'], 'formulas:', sheet['formulas'], 'uncached:', sheet['uncached_formulas'],
                  'errors:', dict(sheet['error_counts']), 'functions:', list(sheet['functions'])[:7])
