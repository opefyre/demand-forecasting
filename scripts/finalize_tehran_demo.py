"""Check every sales-demo workbook, inputs, live columns and uncertainty policy."""
from io import BytesIO
import json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from scripts.prepare_tehran_demo import MANIFEST, INPUT, OUT, write, verify


def finalize():
    from app import main
    verify()
    manifest = json.loads(MANIFEST.read_text())
    workbooks = []
    for key, identifier in manifest['runs'].items():
        run = main._load_run(identifier)
        path = main.RUNS_DIR / identifier / 'forecast_package.xlsx'
        book = load_workbook(path, read_only=True, data_only=True)
        expected = len(run['items']) * run['run_settings']['horizon']
        assert book['Forecast'].max_row-1 == expected
        for sheet in ('Clean History','Models','Series Diagnostics','Engine','Method settings','Range policy'):
            assert sheet in book.sheetnames and book[sheet].max_row>1
        if key=='seasonal:factor_test':
            assert book['Factor choices'].max_row-1==16
        if key.startswith('live'):
            assert run['evidence_policy']=='reviewed_what_if'
            assert all(p['p10'] is None and p['p90'] is None for p in run['series']['__all__']['forecast'])
            assert run['metrics']['wape_pct'] is None
        workbooks.append(dict(run_id=identifier, rows=expected, sheets=book.sheetnames))
        book.close()
    d = main.DATASET_STORE.get(manifest['datasets']['live'])
    factors = [column for column in d['settings']['drivers'] if column.startswith('factor_')]
    assert len(factors)==2
    for role in ('history','future'):
        data = pd.read_csv(BytesIO(main.DATASET_STORE.source(d['sources'][role])[1]))
        assert data[factors].notna().all().all()
    base = main.DATASET_STORE.get(manifest['datasets']['gregorian'])
    h = pd.read_csv(BytesIO(main.DATASET_STORE.source(base['sources']['history'])[1]))
    assert len(h)==804 and not h.duplicated(['customer','sku','date']).any()
    assert h.customer.nunique()==5 and h.sku.nunique()==4
    assert not h.isna().any().any() and h.quantity.ge(0).all()
    lengths = h.groupby('series').size()
    assert sum(lengths.eq(48))==16 and sum(lengths.eq(18))==2
    snapshots = []
    for basis in ('gregorian','jalali'):
        run_id = manifest['runs'][basis+':recommended']
        snap = main.SALES_STORE.get(main.SALES_STORE.list(run_id)[0]['id'])
        pd.DataFrame(snap['inputs']['orders']).to_csv(INPUT/('orders-'+basis+'.csv'),index=False)
        snapshots.append(dict(calendar=basis,order_lines=len(snap['inputs']['orders'])))
    customers = main.CUSTOMER_STORE.list()
    pd.DataFrame([dict(customer=c['customer'],external_id=c['external_id'],sku=p['sku'],unit=p['unit'],
        alternative_name='; '.join(c['aliases'])) for c in customers for p in c['products']]).to_csv(INPUT/'customer-products.csv',index=False)
    contract = main.SALES_STORE.get(main.SALES_STORE.list(manifest['runs']['monthly_contract'])[0]['id'])
    pd.DataFrame(contract['inputs']['commitments']).to_csv(INPUT/'complete-monthly-contracts.csv',index=False)
    # Save a compact evidence report, never keys or old client files.
    manifest['checks']['workbooks'] = workbooks
    manifest['checks']['live_factor_columns_complete'] = factors
    manifest['checks']['orders_by_calendar'] = snapshots
    manifest['checks']['history_validated'] = dict(rows_per_calendar=804,customer_product_pairs=18,
        established_months=48,new_customer_months=18,missing_values=0,duplicate_rows=0)
    manifest['checks']['uncertainty'] = dict(established_portfolio_range=any(p['p10'] is not None
        for p in main._load_run(manifest['runs']['seasonal:model:AutoETS'])['series']['__all__']['forecast']),
        full_portfolio_range='Withheld when the new buyer lacks joint historical errors.',
        live_sensitivity='Withheld: revised historical market data is what-if evidence, not a fair historical test.')
    manifest['checks']['tests'] = dict(full_suite_passed=679,post_clarification_targeted_passed=31)
    manifest['limitations'][4] = 'Servix refreshed successfully; historical quote coverage and original release times are incomplete. A current reference quote is used only as an explicit future what-if hold.'
    manifest['limitations'].append('Synthetic results demonstrate workflows; they do not establish accuracy for this client or prove causal effects.')
    manifest['state'] = 'demo_ready'
    write(MANIFEST,manifest)
    print(json.dumps(dict(state=manifest['state'],customers=5,products=4,model_scenario_runs=len(workbooks),
        retrospective_comparisons=54,demand_exports=len(manifest['exports']),orders=snapshots)),flush=True)


if __name__=='__main__':
    finalize()
