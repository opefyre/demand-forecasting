"""Exercise reviewed factor changes against the existing synthetic plant sample."""
import argparse
from copy import deepcopy
from io import BytesIO
import json
import math
import uuid
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.datasets import DatasetStore
from app.operations import read_operations_workbook

import httpx
import openpyxl

BASE='http://127.0.0.1:8010'
def api(path,payload=None):
    response=httpx.get(BASE+path,timeout=30) if payload is None else httpx.post(BASE+path,json=payload,timeout=30)
    if response.is_error: raise RuntimeError(response.text)
    return response.json()

def start():
    source=next(d for d in api('/api/datasets')['datasets'] if d['id']=='8cfb2da761a840bc8b4404ea7b337773')
    settings=deepcopy(source['settings'])
    settings.update(drivers=['usd_irr_synthetic','energy_curtailment_hours_synthetic'],method_selection='model:Ridge + drivers')
    dataset=api('/api/datasets',{'name':'Synthetic FX and energy baseline','sources':source['sources'],'settings':settings,'classification':'synthetic_sample','accept_warnings':True})
    print(json.dumps(api('/api/jobs',{'dataset_id':dataset['id'],'request_id':str(uuid.uuid4())})))

def verify(identifier):
    scenario=api('/api/runs/'+identifier)
    base=api('/api/runs/'+scenario['base_run_id'])
    assert scenario['scenario']['type']=='factor_assumptions'
    assert scenario['source_classification']==base['source_classification']=='synthetic_sample'
    assert scenario['metrics']['evaluation_signature']==base['metrics']['evaluation_signature']
    assert scenario['metrics']==base['metrics']
    assert scenario['series_ensemble_weights']==base['series_ensemble_weights']
    assert scenario['method_selection']==base['method_selection']
    for item in base['items']:
        assert scenario['series'][item]['history']==base['series'][item]['history']
        changed_periods={row['period'] for row in scenario['scenario']['changes'] if row['item_id']==item}
        for old,new in zip(base['series'][item]['forecast'],scenario['series'][item]['forecast']):
            if str(old['timestamp'])[:10] not in changed_periods:
                assert old['mean']==new['mean'], (item,old['timestamp'])
    original={s['role']:s for s in base['input_manifest']['sources']}
    changed={s['role']:s for s in scenario['input_manifest']['sources']}
    for role in ('history','operations'): assert original[role]==changed[role]
    assert original['future']['sha256']!=changed['future']['sha256']
    # Independently sum recipe requirements; do not call the supply calculator.
    store = DatasetStore(Path(__file__).resolve().parents[1] / 'data/datasets')
    source, content = store.source(changed['operations']['id'])
    tables = read_operations_workbook(source['name'], content, scenario['input_manifest']['settings']['operations_mapping'])
    requirements = {}
    for item in scenario['items']:
        sku = scenario['metadata'][item]['sku']
        for point in scenario['series'][item]['forecast']:
            period = point['timestamp'][:10]
            for recipe in tables['bom'].itertuples():
                if recipe.sku == sku:
                    key = (recipe.material_id, period)
                    requirements.setdefault(key, []).append(point['mean'] * float(recipe.quantity_per_tonne) * (1 + float(recipe.scrap_pct) / 100))
    for row in scenario['operations']['materials']:
        assert row['gross_requirement'] == round(math.fsum(requirements[(row['material_id'], row['period'])]), 3)
    capacities = scenario['operations']['capacity']
    for row in capacities:
        independently_required = math.fsum(p['mean'] for item in scenario['items']
            if scenario['metadata'][item]['production_line'] == row['production_line']
            for p in scenario['series'][item]['forecast'] if p['timestamp'][:10] == row['period'])
        assert row['required_quantity'] == round(independently_required, 3)
    before=math.fsum(row['mean'] for row in base['series']['__all__']['forecast'])
    after=math.fsum(row['mean'] for row in scenario['series']['__all__']['forecast'])
    assert abs(after-before)>0.001
    response=httpx.get(BASE+'/api/export/'+identifier+'/xlsx');response.raise_for_status()
    book=openpyxl.load_workbook(BytesIO(response.content),data_only=True)
    assert book['Changed assumptions'].max_row==len(scenario['scenario']['changes'])+1
    assert book['Factor sources'].max_row==len(scenario['scenario']['definitions'])+1
    print(json.dumps({'base_run_id':base['run_id'],'scenario_run_id':identifier,'base_total':before,'scenario_total':after,'change_pct':(after/before-1)*100,'changed_values':len(scenario['scenario']['changes']),'unchanged_history_tests_and_method':True,'export_verified':True,'material_periods_reconciled':len(requirements),'capacity_periods_reconciled':len(capacities)}))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--start',action='store_true');parser.add_argument('--verify');args=parser.parse_args()
    if args.start:start()
    elif args.verify:verify(args.verify)
    else:parser.error('Choose --start or --verify RUN_ID')
