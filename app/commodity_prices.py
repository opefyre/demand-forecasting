"""World Bank published monthly prices, downloaded directly (not user imports)."""
from datetime import datetime
import io
import math
import re
from urllib.parse import urlsplit
import zipfile

import openpyxl
import pandas as pd

LANDING = 'https://www.worldbank.org/en/research/commodity-markets'
# Reference prices, not local purchasing prices or assumed sales-demand effects.
SERIES = {
    'brent': ('Crude oil, Brent', '($/bbl)'),
    'gas_europe': ('Natural gas, Europe', '($/mmbtu)'),
    'gas_us': ('Natural gas, US', '($/mmbtu)'),
    'aluminum': ('Aluminum', '($/mt)'),
    'copper': ('Copper', '($/mt)'),
    'lead': ('Lead', '($/mt)'),
    'zinc': ('Zinc', '($/mt)'),
    'nickel': ('Nickel', '($/mt)'),
    'iron_ore': ('Iron ore, cfr spot', '($/dmtu)'),
    'cotton': ('Cotton, A Index', '($/kg)'),
    'rubber': ('Rubber, RSS3', '($/kg)'),
    'wheat': ('Wheat, US HRW', '($/mt)'),
    'maize': ('Maize', '($/mt)'),
    'sugar': ('Sugar, world', '($/kg)'),
    'urea': ('Urea', '($/mt)'),
}


def workbook_url(html):
    candidates = re.findall(r'https://[^\s"<>]+/CMO-Historical-Data-Monthly\.xlsx', html)
    urls = sorted({u for u in candidates if urlsplit(u).hostname == 'thedocs.worldbank.org'
                   and urlsplit(u).scheme == 'https' and not urlsplit(u).username
                   and not urlsplit(u).port})
    if len(urls) != 1:
        raise ValueError('The World Bank download link changed. The last saved data is kept.')
    return urls[0]


def parse_prices(content, captured):
    # Bound decompression as well as the network download. Never evaluate formulas.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if len(archive.infolist()) > 500 or sum(r.file_size for r in archive.infolist()) > 30_000_000:
            raise ValueError('The commodity workbook exceeds the expected size.')
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = book['Monthly Prices']
        if sheet.max_row > 2500 or sheet.max_column > 200:
            raise ValueError('The commodity workbook layout changed.')
        rows = list(sheet.values)
        headers = next((i for i, row in enumerate(rows[:20]) if 'Crude oil, Brent' in row), None)
        if headers is None:
            raise ValueError('The commodity price headers changed.')
        result = {}
        cutoff = pd.Timestamp(captured).tz_convert('UTC').date()
        updated = next((str(row[0]) for row in rows[:headers] if str(row[0]).startswith('Updated on')), None)
        labels = [str(x).strip() if x is not None else '' for x in rows[headers]]
        for key, (name, unit) in SERIES.items():
            if labels.count(name) != 1:
                raise ValueError('A commodity price column is missing or duplicated.')
            column = labels.index(name)
            if str(rows[headers + 1][column]).strip() != unit:
                raise ValueError('A commodity price unit changed. No new values were saved.')
            points, seen = [], set()
            for row in rows[headers + 2:]:
                if not row[0]:
                    continue
                match = re.fullmatch(r'(\d{4})M(\d{2})', str(row[0]).strip())
                if not match:
                    raise ValueError('A commodity month is invalid.')
                year, month = map(int, match.groups())
                if not 1960 <= year <= cutoff.year:
                    raise ValueError('A commodity year is invalid.')
                period = (pd.Timestamp(datetime(year, month, 1)) + pd.offsets.MonthEnd()).date()
                if period > cutoff or period in seen:
                    raise ValueError('A commodity month is future-dated or duplicated.')
                seen.add(period)
                value = row[column]
                if value in (None, '', '…', '...'):
                    continue  # Missing observations stay missing; no interpolation.
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                    raise ValueError('A commodity price is invalid.')
                points.append({'period': str(period), 'value': float(value), 'available_at': captured})
            if not points:
                raise ValueError('No commodity observations were returned.')
            result[key] = {'name': name, 'unit': unit.strip('()'), 'points': sorted(points, key=lambda p: p['period']),
                           'provider_updated_at': updated}
        return result
    finally:
        book.close()
