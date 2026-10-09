"""Reviewed, unit-preserving inventory snapshots; never a demand-history source."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import json
import hashlib
import math
from pathlib import Path
import uuid
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd
from sqlalchemy import Column, JSON, MetaData, String, Table, create_engine, select
from sqlalchemy.exc import IntegrityError
from .units import convert_stock


def text_value(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def raw_table(name, content, sheet=None, header_row=1, preview=False):
    try:
        return _raw_table(name, content, sheet, header_row, preview)
    except (BadZipFile, UnicodeDecodeError, pd.errors.ParserError, pd.errors.EmptyDataError) as exc:
        raise ValueError('This file could not be read. Export a valid Excel or UTF-8 CSV stock table and try again.') from exc


def _raw_table(name, content, sheet=None, header_row=1, preview=False):
    """Use positional column IDs so duplicate labels cannot overwrite a measure."""
    if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= 200:
        raise ValueError('Choose the heading row, from 1 to 200.')
    suffix = Path(name).suffix.lower()
    formulas = {}
    if suffix in {'.xlsx', '.xlsm'}:
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheets = book.sheetnames
            sheet = sheet or sheets[0]
            if sheet not in sheets:
                raise ValueError('Choose a worksheet from this file.')
            ws = book[sheet]
            # Valid exports may omit the optional worksheet dimension element.
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)
            if not preview and (ws.max_row * ws.max_column > 1_000_000):
                raise ValueError('This worksheet is too large for inventory import. Export the stock table separately.')
            if ws.max_column > 500:
                raise ValueError('Limit the stock table to 500 columns.')
            last_row = min(ws.max_row, header_row + 8) if preview else ws.max_row
            matrix = [list(row) for row in ws.iter_rows(min_row=header_row, max_row=last_row, values_only=True)]
            total_rows = max(0, ws.max_row - header_row)
        finally:
            book.close()
        if not preview:
            book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                for row in book[sheet].iter_rows(min_row=header_row + 1):
                    for cell in row:
                        if cell.data_type == 'f':
                            formulas[cell.coordinate] = cell.value
            finally:
                book.close()
    elif suffix in {'.csv', '.tsv'}:
        frame = pd.read_csv(BytesIO(content), header=None, dtype=str, keep_default_na=False,
                            sep='\t' if suffix == '.tsv' else ',', encoding='utf-8-sig')
        if frame.size > 1_000_000 or len(frame.columns) > 500:
            raise ValueError('Limit the stock table to 500 columns and 1,000,000 cells.')
        sheets, sheet = [], None
        matrix = frame.iloc[header_row - 1:].values.tolist()
        total_rows = max(0, len(matrix) - 1)
        if preview:
            matrix = matrix[:9]
    elif suffix == '.json':
        records = json.loads(content)
        if not isinstance(records, list) or not all(isinstance(row, dict) for row in records):
            raise ValueError('JSON inventory must be a list of records.')
        if header_row != 1:
            raise ValueError('JSON inventory uses heading row 1.')
        frame = pd.DataFrame(records).where(pd.notna(pd.DataFrame(records)), None)
        if frame.size > 1_000_000 or len(frame.columns) > 500:
            raise ValueError('Limit the stock table to 500 columns and 1,000,000 cells.')
        sheets, sheet = [], None
        matrix = [frame.columns.tolist()] + frame.values.tolist()
        total_rows = len(frame)
        if preview:
            matrix = matrix[:9]
    else:
        raise ValueError('Use Excel (.xlsx), CSV, TSV or JSON for inventory.')
    if not matrix or not any(text_value(value) for value in matrix[0]):
        if preview:
            return {'sheets': sheets, 'sheet': sheet, 'header_row': header_row, 'columns': [],
                    'rows': [], 'source_rows': total_rows, 'formulas': {}}
        raise ValueError('The selected heading row is empty.')
    columns = [{'id': get_column_letter(i + 1), 'label': text_value(value) or f'Column {get_column_letter(i + 1)}'}
               for i, value in enumerate(matrix[0])]
    rows = []
    for number, values in enumerate(matrix[1:], header_row + 1):
        if not any(text_value(value) for value in values):
            continue
        rows.append({'source_row': number, 'values': {col['id']: values[i] if i < len(values) else None
                                                     for i, col in enumerate(columns)}})
    return {'sheets': sheets, 'sheet': sheet, 'header_row': header_row, 'columns': columns,
            'rows': rows, 'source_rows': total_rows, 'formulas': formulas}


def inventory_preview(name, content, sheet=None, header_row=1):
    table = raw_table(name, content, sheet, header_row, preview=True)
    table['rows'] = [{'source_row': row['source_row'], 'values': {key: text_value(value)
                     for key, value in row['values'].items()}} for row in table['rows']]
    table.pop('formulas')
    return table


def review_inventory(table, config):
    if config.get('classification', 'user_provided') not in {'user_provided', 'synthetic_sample'}:
        raise ValueError('Identify the stock as real data or a sample.')
    mapping = config.get('mapping', {})
    columns = {col['id'] for col in table['columns']}
    for field in ('sku', 'quantity'):
        if mapping.get(field) not in columns:
            raise ValueError('Match the product code and on-hand quantity columns.')
    selected = [value for value in mapping.values() if value]
    if len(selected) != len(set(selected)) or any(value not in columns for value in selected):
        raise ValueError('Use a different existing column for each field.')
    mode = config.get('quality_mode', 'unknown')
    if mode not in {'unknown', 'available', 'column'}:
        raise ValueError('Choose how to interpret stock status.')
    if mode == 'column' and mapping.get('quality') not in columns:
        raise ValueError('Choose the stock-status column.')
    fixed_unit = str(config.get('fixed_unit', '')).strip()
    if not mapping.get('unit') and fixed_unit in {'', '-', '?'}:
        raise ValueError('Choose a unit column or enter the confirmed stock unit.')
    try:
        as_of = date.fromisoformat(config.get('as_of', ''))
    except (ValueError, TypeError):
        raise ValueError('Enter the confirmed inventory snapshot date. Packaging dates are not snapshot dates.')
    quality_map = config.get('quality_map', {})
    if any(value not in {'available', 'hold', 'unknown'} for value in quality_map.values()):
        raise ValueError('Stock status must be usable, on hold, or not confirmed.')
    rows, issues, warnings, identities = [], [], [], set()
    repeated, formula_count = 0, 0
    signatures = set()
    for source in table['rows']:
        number, values = source['source_row'], source['values']
        value = lambda field: values.get(mapping.get(field))
        sku, unit = text_value(value('sku')), text_value(value('unit')) if mapping.get('unit') else fixed_unit
        quantity_cell = f'{mapping["quantity"]}{number}'
        quantity = value('quantity')
        try:
            if isinstance(quantity, bool):
                raise ValueError()
            quantity = float(quantity)
            if not math.isfinite(quantity) or quantity < 0:
                raise ValueError()
        except (ValueError, TypeError):
            issues.append({'row': number, 'cell': quantity_cell, 'message': 'Enter a nonnegative on-hand quantity. Missing values and formula errors are not zero.'})
            continue
        if not sku or sku.startswith('#') or unit in {'', '-', '?'} or unit.startswith('#'):
            issues.append({'row': number, 'cell': f'{mapping["sku"]}{number}', 'message': 'Product code or stock unit is missing or invalid.'})
            continue
        keys = [text_value(value(field)) for field in ('record_key', 'record_key_2') if mapping.get(field)]
        if keys:
            identity = (sku, *keys)
            if any(not key for key in keys) or identity in identities:
                issues.append({'row': number, 'cell': f'{mapping.get("record_key", mapping["sku"])}{number}', 'message': 'Stock record keys are blank or repeated for this product. Choose keys that identify separate stock records.'})
                continue
            identities.add(identity)
        raw_quality = text_value(value('quality'))
        quality = quality_map.get(raw_quality, 'unknown') if mode == 'column' else mode
        row = {'sku': sku, 'quantity': quantity, 'unit': unit,
               'location': text_value(value('location')) or 'Unspecified', 'quality': quality,
               'source_status': raw_quality, 'record_keys': keys, 'source_row': number,
               'quantity_cell': quantity_cell, 'sku_cell': f'{mapping["sku"]}{number}'}
        signature = (sku, quantity, unit, row['location'], quality)
        if not keys and signature in signatures:
            repeated += 1
        signatures.add(signature)
        formula_count += int(quantity_cell in table['formulas'])
        rows.append(row)
    if not rows and not issues:
        issues.append({'row': None, 'cell': '', 'message': 'No stock records found below the selected heading row.'})
    if repeated:
        warnings.append(f'{repeated} repeated-looking stock records. Choose record keys to distinguish pallets or lots; no rows were removed.')
    if formula_count:
        warnings.append(f'{formula_count} quantities use saved Excel formula results. Confirm the workbook was recalculated before export.')
    unconfirmed = sum(row['quality'] == 'unknown' and row['quantity'] > 0 for row in rows)
    if unconfirmed:
        warnings.append(f'{unconfirmed} records have unconfirmed stock status. Their products cannot yet have a usable-stock projection.')
    totals = defaultdict(float)
    for row in rows:
        totals[row['unit']] += row['quantity']
    return {'as_of': as_of.isoformat(), 'rows': rows, 'row_count': len(rows),
            'sku_count': len({row['sku'] for row in rows}), 'issues': issues[:100],
            'issue_count': len(issues), 'warnings': warnings, 'totals_by_unit': dict(totals),
            'unknown_quality_rows': unconfirmed, 'duplicate_looking_rows': repeated}


class InventoryStore:
    def __init__(self, path, sources):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.sources = sources
        self.engine = create_engine(f'sqlite:///{path}')
        metadata = MetaData()
        self.table = Table('inventory_snapshots', metadata, Column('id', String, primary_key=True),
                           Column('created_at', String), Column('payload', JSON))
        self.receipts = Table('inventory_receipts', metadata, Column('id', String, primary_key=True),
                              Column('snapshot_id', String), Column('payload', JSON))
        metadata.create_all(self.engine)

    def receipt_versions(self, snapshot_id):
        self.get(snapshot_id)
        with self.engine.connect() as conn:
            rows = list(conn.execute(select(self.receipts.c.payload).where(
                self.receipts.c.snapshot_id == snapshot_id)).scalars())
        return sorted(rows, key=lambda row: row['created_at'], reverse=True)

    def receipt_version(self, snapshot_id, version_id):
        with self.engine.connect() as conn:
            row = conn.execute(select(self.receipts.c.payload).where(
                self.receipts.c.snapshot_id == snapshot_id, self.receipts.c.id == version_id)).scalar()
        if row is None:
            raise ValueError('Receipt schedule not found for this stock snapshot.')
        return row

    def prepare_receipts(self, snapshot_id, config, principal=None):
        snapshot = self.get(snapshot_id)
        imported = None
        if config.get('import_config'):
            from .receipt_imports import import_receipts
            imported = import_receipts(self.sources, config['import_config'])
            config = {**config, 'rows': imported['rows']}
        rows = config.get('rows')
        if not isinstance(rows, list) or len(rows) > 5000:
            raise ValueError('Provide up to 5,000 receipt lines.')
        if config.get('reviewed') is not True:
            raise ValueError('Confirm that quantities are outstanding usable receipts, not already included in stock.')
        name = str(config.get('name', '')).strip()
        reason = str(config.get('reason', '')).strip()
        request_id = str(config.get('request_id', ''))
        if not name or len(name) > 120 or not reason or len(reason) > 1000 or not 1 <= len(request_id) <= 200:
            raise ValueError('Give the receipt schedule a name, a review note and a save request identifier.')
        parent_id = config.get('parent_id')
        if parent_id:
            self.receipt_version(snapshot_id, parent_id)
        cleaned, keys = [], set()
        for i, row in enumerate(rows, 1):
            if not isinstance(row, dict):
                raise ValueError(f'Receipt line {i} is invalid.')
            fields = {key: text_value(row.get(key)) for key in ('reference', 'sku', 'unit', 'due_date', 'kind', 'status')}
            if any(not value or len(value) > 200 for value in fields.values()):
                raise ValueError(f'Receipt line {i}: provide the reference, product, unit, date, type and status.')
            if fields['kind'] not in {'purchase', 'production'} or fields['status'] not in {'confirmed', 'unconfirmed', 'cancelled'}:
                raise ValueError(f'Receipt line {i}: choose a valid type and status.')
            key = fields['reference']
            if key in keys:
                raise ValueError('Each receipt reference must identify one unique order line or scheduled delivery.')
            keys.add(key)
            try:
                due = date.fromisoformat(fields['due_date'])
                if due.isoformat() != fields['due_date'] or due <= date.fromisoformat(snapshot['as_of']):
                    raise ValueError()
                quantity = row.get('quantity')
                if isinstance(quantity, bool) or not math.isfinite(float(quantity)) or float(quantity) <= 0:
                    raise ValueError()
            except (ValueError, TypeError):
                raise ValueError(f'Receipt line {i}: use a positive quantity and a date after the stock snapshot.')
            cleaned.append({**fields, 'quantity': float(quantity), 'source_row': i})
        actor = principal or {'name': 'Local session', 'basis': 'local_session'}
        payload = dict(snapshot_id=snapshot_id, name=name, reason=reason, rows=cleaned,
                       parent_id=parent_id, classification=snapshot.get('classification', 'user_provided'), actor=actor)
        if imported:
            payload['import_evidence'] = {key: value for key, value in imported.items() if key != 'rows'}
        return payload

    def save_receipts(self, snapshot_id, config, principal=None):
        payload = self.prepare_receipts(snapshot_id, config, principal)
        fingerprint = hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()
        key = uuid.uuid5(uuid.NAMESPACE_URL, 'inventory-receipts:' + str(config['request_id'])).hex
        saved = dict(payload, id=key, fingerprint=fingerprint, created_at=datetime.now(timezone.utc).isoformat())
        try:
            with self.engine.begin() as conn:
                conn.execute(self.receipts.insert().values(id=key, snapshot_id=snapshot_id, payload=saved))
        except IntegrityError:
            prior = self.receipt_version(snapshot_id, key)
            if prior['fingerprint'] != fingerprint:
                raise ValueError('This save request already belongs to another receipt schedule.')
            return prior
        return saved

    def inspect(self, config):
        source, content = self.sources.source(config['source_id'])
        if source['role'] != 'inventory':
            raise ValueError('Choose an inventory source, not a demand-history file.')
        table = raw_table(source['name'], content, config.get('sheet'), config.get('header_row', 1))
        review = review_inventory(table, config)
        review['source'] = {key: source[key] for key in ('id', 'name', 'sha256')}
        review['source'].update(sheet=table['sheet'], header_row=table['header_row'])
        return review

    def save(self, config):
        review = self.inspect(config)
        if review['issue_count']:
            raise ValueError('Resolve the listed inventory issues before saving.')
        if not config.get('reviewed'):
            raise ValueError('Confirm the stock date, units and mapped quantities before saving.')
        if review['duplicate_looking_rows'] and not config.get('distinct_records_confirmed'):
            raise ValueError('Choose stock record keys or confirm that repeated-looking rows are distinct physical stock.')
        row = {'id': uuid.uuid4().hex, 'name': str(config.get('name', '')).strip() or 'Inventory snapshot',
               'classification': config.get('classification', 'user_provided'),
               'created_at': datetime.now(timezone.utc).isoformat(), 'config': config, **review}
        with self.engine.begin() as conn:
            conn.execute(self.table.insert().values(id=row['id'], created_at=row['created_at'], payload=row))
        return row

    def get(self, key):
        with self.engine.connect() as conn:
            row = conn.execute(select(self.table.c.payload).where(self.table.c.id == key)).scalar()
        if row is None:
            raise ValueError('Inventory snapshot not found.')
        return row

    def list(self):
        with self.engine.connect() as conn:
            snapshots = list(conn.execute(select(self.table.c.payload).order_by(self.table.c.created_at.desc())).scalars())
        return [{key: value for key, value in snapshot.items() if key not in {'rows', 'config'}} for snapshot in snapshots]


def project_inventory(run, snapshot, unit_version=None, receipt_version=None):
    if unit_version and any(unit_version['classification'] != kind for kind in
                            (snapshot.get('classification', 'user_provided'), run.get('source_classification', 'user_provided'))):
        raise ValueError('Unit definitions, stock and forecast must all be real data or all be samples.')
    if snapshot.get('classification', 'user_provided') != run.get('source_classification', 'user_provided'):
        raise ValueError('Stock and forecast must both be real data or both be samples.')
    demand = defaultdict(lambda: defaultdict(float))
    for item, series in run.get('series', {}).items():
        if item == '__all__':
            continue
        sku = text_value(run.get('metadata', {}).get(item, {}).get('sku') or item)
        for row in series.get('forecast', []):
            quantity = float(row['mean'])
            if not math.isfinite(quantity) or quantity < 0:
                raise ValueError('The forecast contains an invalid quantity. Recalculate it before projecting stock.')
            demand[sku][str(row['timestamp'])[:10]] += quantity
    periods = sorted({period for values in demand.values() for period in values})
    if not periods:
        raise ValueError('This forecast has no future quantities.')
    expected = date.fromisoformat(periods[0]) - timedelta(days=1)
    if snapshot['as_of'] != expected.isoformat():
        raise ValueError(f'This forecast starts {periods[0]}. Use closing stock as of {expected.isoformat()}, or create a forecast beginning immediately after your stock date. Partial-period demand is not guessed.')
    unit = run.get('unit')
    if not unit:
        raise ValueError('Confirm the forecast quantity unit first.')
    receipts = defaultdict(lambda: defaultdict(float))
    receipt_evidence, excluded_receipts = [], []
    if receipt_version:
        if receipt_version['snapshot_id'] != snapshot['id'] or receipt_version['classification'] != snapshot.get('classification', 'user_provided'):
            raise ValueError('Choose a receipt schedule reviewed for this stock snapshot.')
        frequency = run.get('run_settings', {}).get('frequency')
        offsets = {'monthly': pd.DateOffset(months=1), 'weekly': pd.DateOffset(days=7), 'daily': pd.DateOffset(days=1)}
        if frequency not in offsets:
            raise ValueError('Confirm the forecast frequency before including receipts.')
        end = (pd.Timestamp(periods[-1]) + offsets[frequency]).date().isoformat()
        for row in receipt_version['rows']:
            excluded = None
            if row['status'] != 'confirmed': excluded = row['status'].capitalize()
            elif row['sku'] not in demand: excluded = 'Product not in forecast'
            elif not periods[0] <= row['due_date'] < end: excluded = 'Outside forecast dates'
            if excluded:
                excluded_receipts.append({**row, 'exclusion': excluded})
                continue
            period = max(p for p in periods if p <= row['due_date'])
            # Gaps in a forecast must not silently absorb deliveries from absent periods.
            if row['due_date'] >= (pd.Timestamp(period) + offsets[frequency]).date().isoformat() or period not in demand[row['sku']]:
                raise ValueError(f"{row['reference']}: no matching demand period for this receipt.")
            converted, evidence = convert_stock({**row, 'quality': 'available'}, unit, row['due_date'], unit_version)
            receipts[row['sku']][period] += converted['quantity']
            if not math.isfinite(receipts[row['sku']][period]):
                raise ValueError('The combined receipt quantity is too large.')
            receipt_evidence.append({**evidence, 'reference': row['reference'], 'due_date': row['due_date'], 'period': period, 'kind': row['kind']})
    by_sku = defaultdict(list)
    for row in snapshot['rows']:
        by_sku[row['sku']].append(row)
    result, conversions, conversion_issues = [], [], []
    for sku, future in sorted(demand.items()):
        stock = by_sku.get(sku, [])
        reason = None
        if not stock:
            reason = 'Stock not supplied'
        converted = []
        for row in stock:
            try:
                value, evidence = convert_stock(row, unit, snapshot['as_of'], unit_version)
                converted.append(value)
                if row['unit'] != unit:
                    conversions.append(evidence)
            except ValueError as exc:
                reason = 'Unit conversion needed'
                conversion_issues.append({'sku': sku, 'source_row': row['source_row'], 'message': str(exc)})
        if not reason and any(row['quality'] == 'unknown' and row['quantity'] > 0 for row in stock):
            reason = 'Stock status not confirmed'
        available = None if reason else math.fsum(row['quantity'] for row in converted if row['quality'] == 'available')
        held = None if not stock or reason == 'Unit conversion needed' else math.fsum(row['quantity'] for row in converted if row['quality'] == 'hold')
        balance = available
        for period, quantity in sorted(future.items()):
            opening = balance
            received = receipts[sku][period]
            balance = None if balance is None else balance + received - quantity
            if balance is not None and not math.isfinite(balance):
                raise ValueError('The projected stock quantity is too large.')
            result.append({'sku': sku, 'period': period, 'unit': unit, 'opening': opening,
                           'demand': quantity, 'receipts': received, 'held_stock': held, 'closing': balance,
                           'shortfall': None if balance is None else max(0, -balance),
                           'status': reason or ('Shortfall' if balance < 0 else ('Stock remaining' if balance > 0 else 'No stock remaining'))})
    return {'snapshot_id': snapshot['id'], 'as_of': snapshot['as_of'], 'run_id': run['run_id'],
            'stock_classification': snapshot.get('classification', 'user_provided'),
            'forecast_classification': run.get('source_classification', 'user_provided'),
            'plan': run.get('plan'), 'unit': unit, 'rows': result,
            'assumption': ('Period-end balance: opening stock + confirmed outstanding receipts − demand. Negative balances carry forward as unmet demand. Within-period shortages and production feasibility are not determined.' if receipt_version else
                           'On-hand stock only. Future production and purchase receipts are not included. Negative balances carry forward as unmet demand.'),
            'receipt_version': receipt_version, 'receipt_evidence': receipt_evidence, 'excluded_receipts': excluded_receipts,
            'source': snapshot['source'],
            'unit_version': unit_version, 'conversions': conversions, 'conversion_issues': conversion_issues,
            'unmatched_stock_products': sorted(set(by_sku) - set(demand))}
