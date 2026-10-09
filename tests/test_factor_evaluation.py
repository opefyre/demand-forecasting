from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest
import json
import numpy as np
import pandas as pd
from sklearn.linear_model import Ridge
from openpyxl import load_workbook
from app.factor_evaluation import choose_candidate
from app.forecast_engine import ModelSpec,_factor_specs,_fit_backtest,_fit_full_models_and_predict,run_forecast

SPECS=[ModelSpec('Last observed','last'),ModelSpec('Seasonal naive','naive'),ModelSpec('Ridge + drivers','ml',Ridge(alpha=2.))]

def history(periods=84,seed=42):
    rng=np.random.default_rng(seed);dates=pd.date_range('2019-10-01',periods=periods,freq='MS')
    factor=np.repeat(rng.uniform(-2,2,size=(periods+5)//6),6)[:periods]
    noise=rng.normal(0,1,periods)
    rows=[]
    for item in ('Sensitive','Stable'):
        for i,day in enumerate(dates):
            y=180+22*np.sin(i*2*np.pi/12)+(35*factor[i] if item=='Sensitive' else 0)+rng.normal(0,4)
            rows.append({'item_id':item,'timestamp':day,'target':y,'fx':factor[i],'irrelevant':noise[i]})
    return pd.DataFrame(rows)

class FactorEvaluationTests(unittest.TestCase):
    def test_practical_gate_ties_insufficient_history_and_zero_error(self):
        sets={'base':[],'one':['fx'],'both':['fx','noise']}
        self.assertEqual(choose_candidate({'base':1,'one':.97,'both':.98},sets,True)['selected_model'],'base')
        self.assertEqual(choose_candidate({'base':1,'one':.8,'both':.8},sets,True)['selected_model'],'one')
        self.assertEqual(choose_candidate({'base':1,'one':.1,'both':.2},sets,False)['selected_model'],'base')
        self.assertEqual(choose_candidate({'base':0,'one':0,'both':0},sets,True)['selected_model'],'base')
        with self.assertRaisesRegex(ValueError,'history-only'):
            choose_candidate({'base':np.inf,'one':1,'both':1},sets,True)

    def test_bounded_search_uses_existing_models_and_no_exponential_subsets(self):
        with patch('app.forecast_engine._model_specs',return_value=SPECS):
            specs=_factor_specs('deep',['fx','noise'])
        self.assertEqual(len(specs),6)
        self.assertEqual({s.predictors for s in specs if s.kind=='ml'},{(),('fx',),('noise',),('fx','noise')})
        for columns in ([],['fx','fx'],[str(i) for i in range(9)]):
            with self.assertRaises(ValueError):_factor_specs('deep',columns)

    def test_confirmation_does_not_select_factors_and_future_factor_actuals_never_leak(self):
        data=history();changed=data.copy();last=sorted(data.timestamp.unique())[-3:]
        changed.loc[changed.timestamp.isin(last),['fx','irrelevant']]=1e8
        with patch('app.forecast_engine._model_specs',return_value=SPECS):
            original=_fit_backtest(data,['fx','irrelevant'],3,'monthly','deep','factor_test')
            holdout=_fit_backtest(changed,['fx','irrelevant'],3,'monthly','deep','factor_test')
        self.assertEqual(original[3],holdout[3]);self.assertEqual(original[4]['wape_pct'],holdout[4]['wape_pct'])
        changed=data.copy();changed.loc[changed.timestamp.isin(last),'target']*=2
        with patch('app.forecast_engine._model_specs',return_value=SPECS):
            changed_result=_fit_backtest(changed,['fx','irrelevant'],3,'monthly','deep','factor_test')
        self.assertEqual(original[3],changed_result[3])
        self.assertNotEqual(original[4]['wape_pct'],changed_result[4]['wape_pct'])
        report=original[4]['factor_evaluation']
        self.assertFalse(set(report['selection_periods'])&set(report['confirmation_periods']))
        self.assertEqual({r['item_id'] for r in report['rows']},{'Sensitive','Stable'})
        self.assertTrue(all(r['confirmation_points']==3 for r in report['rows']))

    def test_no_factor_variant_truly_ignores_future_values_at_full_refit(self):
        data=history();future=pd.DataFrame([{'item_id':item,'timestamp':day,'fx':100,'irrelevant':-999} for item in ('Sensitive','Stable') for day in pd.date_range('2026-10-01',periods=3,freq='MS')])
        with patch('app.forecast_engine._model_specs',return_value=SPECS):specs=_factor_specs('deep',['fx','irrelevant'])
        before=_fit_full_models_and_predict(data,future,['fx','irrelevant'],'monthly',specs)[0]
        after=_fit_full_models_and_predict(data,future.assign(fx=1e8,irrelevant=1e9),['fx','irrelevant'],'monthly',specs)[0]
        for name in before:
            if '[history only]' in name or '[' not in name:self.assertEqual(before[name],after[name])
        self.assertNotEqual(before['Ridge + drivers [fx]'],after['Ridge + drivers [fx]'])

    def test_short_history_falls_back_and_zero_actuals_do_not_create_accuracy_percentages(self):
        data=history(periods=18).assign(target=0.)
        with patch('app.forecast_engine._model_specs',return_value=SPECS):
            result=_fit_backtest(data,['fx'],6,'monthly','deep','factor_test')
        self.assertTrue(all(not r['selected_factors'] for r in result[4]['factor_evaluation']['rows']))
        self.assertIsNone(result[4]['wape_pct'])

    def test_fast_catalog_still_runs_required_factor_selection_windows(self):
        with patch('app.forecast_engine._model_specs',return_value=SPECS):
            result=_fit_backtest(history(),['fx'],3,'monthly','fast','factor_test')
        report=result[4]['factor_evaluation']
        self.assertGreaterEqual(len(report['selection_periods']),6)
        self.assertEqual(len(report['confirmation_periods']),3)
        self.assertTrue(all(r['reason']!='More complete test windows are needed.' for r in report['rows']))

    def test_real_engine_outputs_report_and_reconciled_workbook_without_mutating_inputs(self):
        data=history();saved=deepcopy(data)
        future=pd.DataFrame([{'item_id':item,'timestamp':day,'fx':.8,'irrelevant':0.} for item in ('Sensitive','Stable') for day in pd.date_range('2026-10-01',periods=3,freq='MS')])
        with TemporaryDirectory() as tmp,patch('app.forecast_engine._model_specs',return_value=SPECS):
            result=run_forecast(history=data,future_covariates=future,known_driver_cols=['fx','irrelevant'],horizon=3,frequency='monthly',profile='deep',runs_dir=Path(tmp),method_selection='factor_test')
            folder=Path(tmp)/result['run_id'];report=json.loads((folder/'factor_evaluation.json').read_text())
            self.assertEqual(report,result['factor_evaluation'])
            workbook=load_workbook(folder/'forecast_package.xlsx',read_only=True);self.addCleanup(workbook.close)
            self.assertIn('Factor choices',workbook.sheetnames);self.assertEqual(workbook['Factor choices'].max_row,3)
            exported=pd.read_csv(folder/'forecast.csv');self.assertEqual(len(exported),6)
            for i,total in enumerate(result['series']['__all__']['forecast']):
                self.assertAlmostEqual(total['mean'],sum(result['series'][key]['forecast'][i]['mean'] for key in ('Sensitive','Stable')))
            self.assertFalse(result['metrics']['factor_release_dates_verified'])
        pd.testing.assert_frame_equal(data,saved)

    def test_retrospective_history_and_missing_factor_values_cannot_use_automatic_choice(self):
        data=history();future=pd.DataFrame([{'item_id':'Sensitive','timestamp':'2026-10-01','fx':None}])
        with TemporaryDirectory() as tmp:
            args=dict(history=data,future_covariates=future,known_driver_cols=['fx'],horizon=1,frequency='monthly',profile='deep',runs_dir=Path(tmp),method_selection='factor_test')
            with self.assertRaisesRegex(ValueError,'what-if'):run_forecast(**args,evidence_policy='reviewed_what_if')
            with self.assertRaisesRegex(ValueError,'missing is not zero'):run_forecast(**args)

    def test_generated_scenarios_multiple_seeds_keep_frozen_choice_complete_and_honest(self):
        from scripts.generate_factor_test_samples import generate,CASES
        for case in CASES:
            for seed in (7301,8107):
                with self.subTest(case=case,seed=seed):
                    raw,_,_,_=generate(case,seed)
                    data=raw.rename(columns={'series':'item_id','date':'timestamp','quantity':'target'})
                    data.timestamp=pd.to_datetime(data.timestamp)
                    with patch('app.forecast_engine._model_specs',return_value=SPECS):
                        result=_fit_backtest(data,['exchange_rate','unrelated_index'],6,'monthly','deep','factor_test')
                    report=result[4]['factor_evaluation']
                    self.assertEqual(len(report['rows']),4)
                    self.assertFalse(set(report['selection_periods'])&set(report['confirmation_periods']))
                    for row in report['rows']:
                        self.assertTrue(np.isfinite(row['selection_selected_loss']))
                        self.assertLessEqual(row['selection_selected_loss'],row['selection_baseline_loss'])
                        if row['selected_factors']:self.assertGreaterEqual(row['selection_gain_pct'],5)
                        if case=='short_history':self.assertEqual(row['selected_factors'],[])
                    self.assertFalse(result[4]['factor_release_dates_verified'])

    def test_real_persian_month_boundaries_are_kept_in_factor_test_exports(self):
        from app.sales_conventions import month_range
        data=history();dates=month_range('2019-10-23','2026-09-23',basis='jalali')
        self.assertEqual(len(dates),84)
        for item in data.item_id.unique():data.loc[data.item_id==item,'timestamp']=dates
        data.attrs['calendar_profile']={'month_basis':'jalali'}
        future=pd.DataFrame([{'item_id':item,'timestamp':day,'fx':.8,'irrelevant':0.} for item in data.item_id.unique()
            for day in month_range('2026-10-23','2026-12-22',basis='jalali')])
        with TemporaryDirectory() as tmp,patch('app.forecast_engine._model_specs',return_value=SPECS):
            result=run_forecast(history=data,future_covariates=future,known_driver_cols=['fx'],horizon=3,frequency='monthly',profile='deep',runs_dir=Path(tmp),method_selection='factor_test')
            frame=pd.read_csv(Path(tmp)/result['run_id']/'forecast.csv')
            self.assertEqual(set(frame.period_label),{'1405-08','1405-09','1405-10'})
            self.assertEqual(set(frame.planning_calendar),{'jalali'})
