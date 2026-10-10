import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import uuid
from unittest.mock import patch

from fastapi.testclient import TestClient
from app.cloud_api import CONTRACT, route
from app.cloud_compute import app
from app.cloud_state import capture, restore
from app.company_workspace import CompanyWorkspaces
from app.platform_api import create_platform_api
from tests import test_platform_sales as sales_fixture


class CloudApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.client = TestClient(app)
        self.who = {'company_id': 'tehran_a', 'issuer': 'https://forecast.vrolen.com', 'subject': 'planner-a',
            'role': 'planner', 'auth_kind': 'session', 'session_id': 'test-only', 'mfa_required': False,
            'permissions': sales_fixture.ALL.copy()}
        self.environment = patch.dict('os.environ', {'DEMANDLAB_CLOUD_RUNTIME': 'true', 'DEMANDLAB_COMPANY_VAULT_KEY': 'ab'*32})
        self.environment.start()
        self.attempts = []

    def tearDown(self):
        for attempt in self.attempts: self.client.delete('/output/' + attempt)
        self.client.close()
        self.environment.stop()
        self.temp.cleanup()

    def command(self, method, path, body, checkpoint=b'', who=None, content_type='application/json'):
        attempt = uuid.uuid4().hex
        self.attempts.append(attempt)
        staged = self.client.put('/request/' + attempt, content=body if isinstance(body,bytes) else json.dumps(body).encode())
        self.assertEqual(staged.status_code, 200, staged.text)
        response = self.client.post('/api-operation', content=checkpoint, headers={'x-forecast-job': json.dumps({
            'company_id': 'tehran_a', 'job_id': attempt, 'payload': {
                'method': method, 'path': path, 'content_type': content_type, 'principal': who or self.who}})})
        if response.status_code != 200: return response, None, None, b''
        metadata = response.json()
        value = self.client.get('/output/' + attempt + '/artifact/api-response.json').json()
        if method == 'GET':
            self.assertEqual(self.client.get('/output/' + attempt + '/snapshot').status_code, 404)
            return response, value, None, checkpoint
        if metadata['api_status'] >= 300:
            self.assertEqual(self.client.get('/output/' + attempt + '/snapshot').status_code, 404)
            return response, value, None, b''
        view = self.client.get('/output/' + attempt + '/artifact/company-view.json').json()
        checkpoint = self.client.get('/output/' + attempt + '/snapshot').content
        self.client.delete('/output/' + attempt)
        return response, value, view, checkpoint

    def test_contract_routes_exist_and_reject_unscoped_or_network_actions(self):
        api = create_platform_api(None, CompanyWorkspaces(self.root/'contract'))
        import re
        actual = {(method.upper(),re.sub(r'\{[^}]+\}', '{id}', path)) for path,methods in api.openapi()['paths'].items() for method in methods}
        for entry in CONTRACT:
            path = entry['path'].replace('{id}', 'test-id').replace('{index}', '0')
            self.assertTrue(any(method == entry['method'] and re.fullmatch(re.escape(pattern).replace(r'\{id\}', '[^/]+'),path)
                for method,pattern in actual),entry)
        for path,methods in api.openapi()['paths'].items():
            # Identity is intentionally native Better Auth/D1, not a scratch API.
            if path.split('/')[1] in {'me','api-keys','access-options','members','invitations','audit-events'}:continue
            examples=[path]
            for parameter,values in [('kind',['csv','xlsx','models','drivers'] if '/files/' in path else ['sources','datasets','forecasts','runs']),
                                     ('role',['customers','orders','commitments'])]:
                if '{'+parameter+'}' in path:
                    examples=[example.replace('{'+parameter+'}',value) for example in examples for value in values]
            for method in methods:
                for example in examples:
                    example=re.sub(r'\{[^}]+\}',lambda match:'0' if match.group(0)=='{index}' else 'test-id',example)
                    self.assertEqual(route(method.upper(),example)['method'],method.upper(),(method,path))
        for path in ['/api-keys', '/members', '/assistant', '/external-sources/refresh', '/../customers', '/customers/a/b']:
            with self.assertRaises(ValueError): route('POST', path)

    def test_four_customers_survive_cold_updates_archiving_and_validation_errors(self):
        checkpoint = b''
        records = []
        for name in ['Mehr Packaging', 'Aftab Printing', 'Pars Trading', 'Negin Cartons']:
            response, customer, view, checkpoint = self.command('POST', '/customers', {
                'customer': name, 'products': [{'sku': '001', 'unit': 'tonnes'}]}, checkpoint)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['api_status'], 201)
            records.append(customer)
            self.assertEqual(view['views']['/customers']['total'], len(records))
        response, _, view, checkpoint = self.command('PUT', '/customers/'+records[0]['id'], {
            'customer':'Mehr Packaging Tehran', 'products':[{'sku':'001','unit':'tonnes'},{'sku':'002','unit':'tonnes'}]}, checkpoint)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(len(view['views']['/customers/'+records[0]['id']+'/products']['products']),2)
        response, _, view, checkpoint = self.command('DELETE', '/customers/'+records[1]['id'], None, checkpoint)
        self.assertEqual(response.status_code,200,response.text)
        self.assertFalse(view['views']['/customers/'+records[1]['id']]['active'])
        unchanged = checkpoint
        response, error, _, bad_checkpoint = self.command('POST', '/customers', {'customer':''}, checkpoint)
        self.assertEqual(response.json()['api_status'],422)
        self.assertTrue(error.get('detail'))
        self.assertEqual(bad_checkpoint,b'')
        wrong = {**self.who,'company_id':'tehran_b'}
        self.assertEqual(self.command('POST','/customers',{'customer':'Wrong company'},checkpoint,wrong)[0].status_code,422)
        self.assertEqual(self.command('POST','/customers',{'customer':'No write'},checkpoint,{**self.who,'permissions':['customers:read']})[0].status_code,422)
        archive=self.root/'restored.zip';archive.write_bytes(unchanged)
        restore(archive,self.root/'restored','tehran_a')
        workspaces=CompanyWorkspaces(self.root/'restored/data/companies')
        try: self.assertEqual(len(workspaces.for_principal(self.who).customers.list()),4)
        finally: workspaces.close()

    def test_real_grouped_models_and_partial_orders_have_saved_read_views(self):
        fixture=sales_fixture.PublicSalesTests();fixture.setUp()
        try:
            _, dataset, _ = fixture.history()
            orders=fixture.orders(dataset)
            import shutil
            shutil.copytree(fixture.root/'companies',self.root/'input/data/companies')
            archive=self.root/'input.zip';capture(self.root/'input','tehran_a',archive)
            response, group, view, checkpoint=self.command('POST','/forecasts',{
                'name':'Tehran sales outlook','dataset_id':dataset['id'],'sales_input_id':orders['id'],
                'methods':['model:Last observed','model:Recent average'],'request_id':'cloud-grouped-forecast'},archive.read_bytes())
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(response.json()['api_status'],202)
            self.assertEqual(len(group['jobs']),2)
            self.assertTrue(all(job['state']=='succeeded' for job in group['jobs']))
            self.assertEqual(view['views']['/forecasts']['total'],1)
            self.assertEqual(view['views']['/runs']['total'],2)
            self.assertIn('/order-snapshots/'+orders['id'],view['views'])
            for job in group['jobs']:
                demand=view['views']['/runs/'+job['run_id']+'/demand']
                current=[r for r in demand['rows'] if r['period']==str(fixture.today.replace(day=1))]
                self.assertEqual(sum(r['total'] for r in current),49)
                self.assertNotIn('job_owner',view['views']['/runs/'+job['run_id']])
            self.assertTrue(checkpoint.startswith(b'PK'))
        finally: fixture.tearDown()

    def test_upload_mapping_order_book_and_review_use_existing_company_routes(self):
        fixture=sales_fixture.PublicSalesTests();fixture.setUp()
        try:
            source, dataset, _=fixture.history()
            _, saved=fixture.ws.datasets.source(source['id'])
            boundary='forecast-cloud-test-boundary'
            body=(f'--{boundary}\r\nContent-Disposition: form-data; name="role"\r\n\r\nhistory\r\n'
                f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="tehran-sales.csv"\r\n'
                f'Content-Type: text/csv\r\n\r\n').encode()+saved+f'\r\n--{boundary}--\r\n'.encode()
            response, uploaded, _, checkpoint=self.command('POST','/sources',body,
                content_type='multipart/form-data; boundary='+boundary)
            self.assertEqual(response.status_code,200,response.text)
            values={'name':'Tehran monthly sales','sources':{'history':uploaded['id']},'settings':dataset['settings'],
                'classification':'synthetic_sample','request_id':'cloud-input-mapping'}
            response, review, _, checkpoint=self.command('POST','/datasets/preview',values,checkpoint)
            self.assertEqual(response.json()['api_status'],200,review)
            response, imported, view, checkpoint=self.command('POST','/datasets',values,checkpoint)
            self.assertEqual(response.json()['api_status'],201,imported)
            starter=view['views']['/datasets/'+imported['id']+'/orders']
            self.assertEqual(len(starter['inputs']['customers']),4)
            from datetime import timedelta
            inputs=starter['inputs'] | {'reviewed':True,'order_feed':'complete_snapshot',
                'valid_until':str(fixture.today+timedelta(days=30)),
                'note':'Reviewed the complete order book: no confirmed orders yet.'}
            path='/datasets/'+imported['id']+'/orders'
            response, book, view, checkpoint=self.command('PUT',path,{
                'version':starter['version'],'as_of':inputs['as_of'],'valid_until':inputs['valid_until'],
                'order_feed':'complete_snapshot','orders':[]},checkpoint)
            self.assertEqual(response.json()['api_status'],200,book)
            values={'inputs':inputs,'request_id':'cloud-orders-review'}
            response, reviewed, _, checkpoint=self.command('POST',path+'/preview',values,checkpoint)
            self.assertEqual(response.json()['api_status'],200,reviewed)
            response, snapshot, view, checkpoint=self.command('POST',path+'/snapshots',values | {'review_token':reviewed['review_token']},checkpoint)
            self.assertEqual(response.json()['api_status'],201,snapshot)
            self.assertIn('/order-snapshots/'+snapshot['id'],view['views'])
        finally: fixture.tearDown()

    def test_approved_reports_independent_review_exports_and_lifecycle_across_cold_starts(self):
        fixture=sales_fixture.PublicSalesTests();fixture.setUp()
        try:
            _,dataset,_=fixture.history();orders=fixture.orders(dataset)
            group,_=fixture.calculate(dataset,orders,['model:Last observed'])
            run_id=group['jobs'][0]['run_id'];run=fixture.ws.load_run(run_id)
            source=self.root/'approved.zip'
            import shutil
            shutil.copytree(fixture.root/'companies',self.root/'input/data/companies')
            capture(self.root/'input','tehran_a',source);checkpoint=source.read_bytes()
            payload={'snapshot_id':run['sales_input_snapshot_id'],'receiver':'Client MRP',
                'mode':'combined_demand','request_id':'cloud-release-review'}
            response,review,view,checkpoint=self.command('POST','/releases/preview',payload,checkpoint)
            self.assertEqual(response.json()['api_status'],200,review)
            self.assertEqual(view['report_views']['/runs']['runs'],[])
            response,record,view,checkpoint=self.command('POST','/releases',payload|{'reviewed':True,'review_token':review['review_token']},checkpoint)
            self.assertEqual(response.json()['api_status'],201,record)
            approval={'reviewed':True,'review_token':review['review_token']}
            own=dict(self.who,role='approver')
            response,_,_,_=self.command('POST','/releases/'+record['id']+'/approve',approval,checkpoint,who=own)
            self.assertEqual(response.json()['api_status'],403)
            approver=dict(own,subject='independent-reviewer')
            response,approved,view,checkpoint=self.command('POST','/releases/'+record['id']+'/approve',approval,checkpoint,who=approver)
            self.assertEqual(response.json()['api_status'],200,approved)
            self.assertEqual(len(view['report_views']['/runs']['runs']),1)
            self.assertIn('/runs/'+run_id,view['report_views'])
            self.assertNotIn('/runs/'+run_id+'/demand',view['report_views'])
            viewer=dict(self.who,subject='reader',role='viewer',permissions=['reports:read','reports:export'])
            response,fresh,read_view,unchanged=self.command('GET','/runs?limit=1000',None,checkpoint,who=viewer)
            self.assertEqual(response.json()['api_status'],200,fresh)
            self.assertEqual(len(fresh['runs']),1)
            self.assertEqual(unchanged,checkpoint)
            self.assertIn('company-view.json',response.json()['artifacts'])
            for kind in ['csv','xlsx','json']:
                response,descriptor,_,unchanged=self.command('GET','/releases/'+record['id']+'/export?kind='+kind,None,checkpoint,who=viewer)
                self.assertEqual(response.json()['api_status'],200,descriptor)
                self.assertEqual(unchanged,checkpoint)
                binary=self.client.get('/output/'+self.attempts[-1]+'/artifact/api-download.bin').content
                self.assertGreater(len(binary),100)
                if kind=='xlsx':self.assertTrue(binary.startswith(b'PK'))
                if kind=='json':self.assertIsInstance(json.loads(binary),list)
            path='/datasets/'+dataset['id']
            response,metadata,view,checkpoint=self.command('PATCH',path+'/metadata',{'name':'Tehran demand inputs','version':0},checkpoint)
            self.assertEqual(response.json()['api_status'],200,metadata)
            self.assertEqual(view['views'][path+'/metadata']['name'],'Tehran demand inputs')
            version=metadata['lifecycle']['version']
            response,archived,view,checkpoint=self.command('POST',path+'/archive',{'version':version},checkpoint)
            self.assertTrue(archived['lifecycle']['archived'])
            response,restored,view,checkpoint=self.command('POST',path+'/restore',{'version':archived['lifecycle']['version']},checkpoint)
            self.assertFalse(restored['lifecycle']['archived'])
            response,_,_,_=self.command('GET','/runs/'+run_id+'/export?mode=combined_demand&kind=csv',None,checkpoint,who=viewer)
            self.assertEqual(response.status_code,422,'Viewer cannot ask for draft exports')
        finally:fixture.tearDown()

    def test_corrected_history_to_revised_group_independent_approval_and_exact_exports(self):
        """Run the full sales journey through durable cold-start operations, offline."""
        import csv
        import shutil
        from io import StringIO, BytesIO
        from openpyxl import load_workbook
        fixture = sales_fixture.PublicSalesTests(); fixture.setUp()
        try:
            source, original, original_body = fixture.history()
            orders = fixture.orders(original)
            parent, _ = fixture.calculate(original, orders, ['model:Last observed'])
            parent_run_id = parent['jobs'][0]['run_id']
            original_run = fixture.ws.load_run(parent_run_id)
            _, original_bytes = fixture.ws.datasets.source(source['id'])
            # Correct a recent customer's demand and remove an old month explicitly.
            rows = list(csv.DictReader(StringIO(original_bytes.decode())))
            last = max(row['date'] for row in rows)
            first = min(row['date'] for row in rows)
            rows = [row for row in rows if row['date'] != first]
            for row in rows:
                if row['date'] == last and row['customer'] == 'Aftab': row['quantity'] = '26'
            output = StringIO(); writer = csv.DictWriter(output, fieldnames=list(rows[0]))
            writer.writeheader(); writer.writerows(rows)
            corrected_bytes = output.getvalue().encode()
            shutil.copytree(fixture.root/'companies', self.root/'input/data/companies')
            archive = self.root/'journey.zip'; capture(self.root/'input', 'tehran_a', archive)
            checkpoint = archive.read_bytes()
            boundary = 'forecast-revision-test'
            upload = (f'--{boundary}\r\nContent-Disposition: form-data; name="role"\r\n\r\nhistory\r\n'
                f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="sales.csv"\r\n'
                'Content-Type: text/csv\r\n\r\n').encode()+corrected_bytes+f'\r\n--{boundary}--\r\n'.encode()
            response, corrected, _, checkpoint = self.command('POST', '/sources', upload, checkpoint,
                content_type='multipart/form-data; boundary='+boundary)
            self.assertEqual(response.json()['api_status'], 201, corrected)
            values = original_body | {'name':'Tehran revised sales history', 'sources':{'history':corrected['id']},
                'parent_dataset_id':original['id'], 'request_id':'corrected-input-journey', 'accept_warnings':False}
            response, review, _, checkpoint = self.command('POST', '/datasets/preview', values, checkpoint)
            self.assertEqual(response.json()['api_status'], 200, review)
            receipt = review['repeat_upload']
            self.assertEqual(receipt['changed_groups'], 1)
            self.assertEqual(receipt['removed_groups'], 4)
            self.assertEqual(receipt['before_rows'], 144)
            self.assertEqual(receipt['after_rows'], 140)
            self.assertEqual(receipt['mode'], 'replacement_not_append')
            response, rejected, _, failed_checkpoint = self.command('POST', '/datasets', values, checkpoint)
            self.assertEqual(response.json()['api_status'], 400, rejected)
            self.assertIn('repeat-upload', rejected['detail'])
            self.assertEqual(failed_checkpoint, b'')
            values['accept_warnings'] = True
            response, revised, view, checkpoint = self.command('POST', '/datasets', values, checkpoint)
            self.assertEqual(response.json()['api_status'], 201, revised)
            self.assertEqual(revised['import_provenance']['review'], receipt)
            self.assertEqual(revised['parent_dataset_id'], original['id'])
            response, retry, _, checkpoint = self.command('POST', '/datasets', values, checkpoint)
            self.assertEqual(retry['id'], revised['id'], 'Retry must not create duplicate input versions')
            # Reuse confirmed orders, but require a new review against the revised input.
            starter = view['views']['/datasets/'+revised['id']+'/orders']
            inputs = starter['inputs'] | {'orders':orders['inputs']['orders'], 'reviewed':True,
                'order_feed':'complete_snapshot', 'valid_until':orders['inputs']['valid_until'],
                'note':'Reviewed existing confirmed orders against corrected sales history.'}
            path = '/datasets/'+revised['id']+'/orders'
            order_body = {'inputs':inputs, 'request_id':'revised-orders-journey'}
            response, order_review, _, checkpoint = self.command('POST', path+'/preview', order_body, checkpoint)
            self.assertEqual(response.json()['api_status'], 200, order_review)
            response, new_orders, _, checkpoint = self.command('POST', path+'/snapshots',
                order_body | {'review_token':order_review['review_token']}, checkpoint)
            self.assertEqual(response.json()['api_status'], 201, new_orders)
            group_body = {'name':'Tehran revised monthly demand', 'dataset_id':revised['id'],
                'sales_input_id':new_orders['id'], 'methods':['model:Last observed','model:Recent average'],
                'parent_forecast_id':parent['id'], 'request_id':'revised-group-journey'}
            response, group, view, checkpoint = self.command('POST', '/forecasts', group_body, checkpoint)
            self.assertEqual(response.json()['api_status'], 202, group)
            self.assertEqual(len(group['jobs']), 2)
            self.assertTrue(all(job['state']=='succeeded' for job in group['jobs']), group)
            self.assertEqual(view['views']['/forecasts']['total'], 2, 'Two methods belong to one revised group')
            versions = view['views']['/forecasts/'+group['id']+'/revisions']
            self.assertEqual({row['id'] for row in versions['revisions']}, {parent['id'],group['id']})
            run_id = group['jobs'][0]['run_id']; run = view['views']['/runs/'+run_id]
            demand = view['views']['/runs/'+run_id+'/demand']
            current = [row for row in demand['rows'] if row['period']==str(fixture.today.replace(day=1))]
            self.assertEqual({row['customer'] for row in current}, {'Mehr','Aftab','Pars','Negin'})
            by_customer = {row['customer']:row for row in current}
            self.assertEqual(by_customer['Aftab']['baseline'], 26)
            self.assertGreater(by_customer['Mehr']['booked'], by_customer['Mehr']['baseline'])
            self.assertEqual(by_customer['Pars']['booked'], 0)
            self.assertEqual(by_customer['Pars']['remaining'], 5)
            self.assertEqual(by_customer['Negin']['remaining'], 8)
            release_body = {'snapshot_id':run['sales_input_snapshot_id'], 'receiver':'Client MRP',
                'mode':'combined_demand', 'request_id':'revised-release-journey'}
            response, report, _, checkpoint = self.command('POST', '/releases/preview', release_body, checkpoint)
            self.assertEqual(response.json()['api_status'], 200, report)
            response, release, _, checkpoint = self.command('POST', '/releases', release_body |
                {'reviewed':True,'review_token':report['review_token']}, checkpoint)
            self.assertEqual(response.json()['api_status'], 201, release)
            export_path = '/releases/'+release['id']+'/export?kind='
            response, _, _, _ = self.command('GET', export_path+'json', None, checkpoint)
            self.assertEqual(response.json()['api_status'], 409, 'Pending quantities cannot be exported as approved')
            approval = {'reviewed':True,'review_token':report['review_token']}
            response, _, _, _ = self.command('POST', '/releases/'+release['id']+'/approve', approval, checkpoint,
                who=dict(self.who,role='approver'))
            self.assertEqual(response.json()['api_status'], 403, 'Creator cannot independently approve')
            response, approved, _, checkpoint = self.command('POST', '/releases/'+release['id']+'/approve', approval,
                checkpoint, who=dict(self.who,role='approver',subject='independent-reviewer'))
            self.assertEqual(response.json()['api_status'], 200, approved)
            viewer = dict(self.who,role='viewer',subject='report-reader',permissions=['reports:read','reports:export'])
            expected = {(row['customer'],row['sku'],row['period']):row['still_to_serve'] for row in demand['rows']}
            for kind in ['json','csv','xlsx']:
                response, descriptor, _, unchanged = self.command('GET', export_path+kind, None, checkpoint, who=viewer)
                self.assertEqual(response.json()['api_status'], 200, descriptor)
                self.assertEqual(unchanged, checkpoint)
                binary = self.client.get('/output/'+self.attempts[-1]+'/artifact/api-download.bin').content
                if kind=='json': exported = json.loads(binary)
                elif kind=='csv': exported = list(csv.DictReader(StringIO(binary.decode('utf-8-sig'))))
                else:
                    workbook = load_workbook(BytesIO(binary),read_only=True)
                    sheet_rows = workbook['Demand'].iter_rows(values_only=True)
                    headers = next(sheet_rows); exported = [dict(zip(headers,row)) for row in sheet_rows]
                    workbook.close()
                self.assertEqual(len(exported),len(expected))
                for row in exported:
                    self.assertIn(row['sku'], {'001','002'}, 'SKU leading zeroes must survive each format')
                    self.assertEqual(row['approval'],'approved')
                    self.assertAlmostEqual(float(row['quantity']),expected[(row['customer'],row['sku'],row['period'])],places=9)
            wrong = dict(viewer,company_id='tehran_b')
            self.assertEqual(self.command('GET',export_path+'json',None,checkpoint,who=wrong)[0].status_code,422)
            restored_archive = self.root/'journey-restored.zip'; restored_archive.write_bytes(checkpoint)
            restore(restored_archive,self.root/'restored','tehran_a')
            workspaces = CompanyWorkspaces(self.root/'restored/data/companies')
            try:
                ws = workspaces.for_principal(self.who)
                self.assertEqual(ws.datasets.source(source['id'])[1],original_bytes)
                self.assertEqual(ws.datasets.get(original['id'])['sources'],original['sources'])
                self.assertEqual(ws.load_run(parent_run_id),original_run)
                self.assertEqual(ws.datasets.get(revised['id'])['import_provenance']['review'],receipt)
            finally: workspaces.close()
        finally: fixture.tearDown()

    def test_grouped_monthly_factor_revision_survives_cold_checkpoint_without_duplicate_jobs(self):
        from copy import deepcopy
        from io import BytesIO
        import pandas as pd
        import shutil
        fixture=sales_fixture.PublicSalesTests();fixture.setUp()
        try:
            _,dataset,_=fixture.history()
            parent,_=fixture.calculate(dataset,fixture.orders(dataset),['model:Last observed'])
            base_id=parent['jobs'][0]['run_id'];original=deepcopy(fixture.ws.load_run(base_id))
            history=pd.read_csv(BytesIO(fixture.ws.datasets.source(dataset['sources']['history'])[1]),dtype={'sku':str})
            history['fx_assumption']=history.groupby('customer').cumcount()+100
            upload=fixture.ws.datasets.upload('tehran-factors.csv',history.to_csv(index=False).encode(),'history')
            future=pd.DataFrame({'date':pd.date_range(fixture.today.replace(day=1),periods=3,freq='MS').strftime('%Y-%m-%d'),
                'fx_assumption':[136,138,140]})
            assumptions=fixture.ws.datasets.upload('reviewed-future.csv',future.to_csv(index=False).encode(),'future')
            adjusted=fixture.ws.datasets.save('Tehran reviewed factors',{'history':upload['id'],'future':assumptions['id']},
                dataset['settings']|{'drivers':['fx_assumption'],'future_date_col':'date','future_driver_policy':'require',
                    'evidence_policy':'reviewed_what_if'},'synthetic_sample',True,parent_dataset_id=dataset['id'])
            orders=fixture.orders(adjusted)
            shutil.copytree(fixture.root/'companies',self.root/'grouped-monthly/data/companies')
            archive=self.root/'grouped-monthly.zip';capture(self.root/'grouped-monthly','tehran_a',archive)
            response,state,_,checkpoint=self.command('POST','/forecast-updates',
                {'run_id':base_id,'workflow':'grouped','request_id':'grouped-cold-start'},archive.read_bytes())
            self.assertEqual(response.json()['api_status'],201,state)
            path='/forecast-updates/'+state['id']
            body={'action':'calculate_group','revision':state['revision'],'reviewed':True,
                'name':'Tehran factor revision','dataset_id':adjusted['id'],'sales_input_id':orders['id'],
                'methods':['model:Ridge + drivers','model:Elastic Net + drivers'],'request_id':'grouped-cold-calculate'}
            response,state,view,checkpoint=self.command('POST',path+'/steps',body,checkpoint)
            self.assertEqual(response.json()['api_status'],200,state)
            self.assertEqual(state['stage'],'grouped_complete');self.assertEqual(len(state['group']['jobs']),2)
            self.assertTrue(all(j['state']=='succeeded' for j in state['group']['jobs']),state)
            self.assertEqual(view['views']['/forecasts']['total'],2)
            response,resumed,_,_=self.command('GET',path,None,checkpoint)
            self.assertEqual(resumed['group']['id'],state['group']['id'])
            response,retried,view,repeated=self.command('POST',path+'/steps',body,checkpoint)
            self.assertEqual(response.json()['api_status'],200,retried)
            self.assertEqual(retried['group']['id'],state['group']['id'])
            self.assertEqual(view['views']['/forecasts']['total'],2)
            for job in resumed['group']['jobs']:
                response,run,_,_=self.command('GET','/runs/'+job['run_id'],None,repeated)
                self.assertEqual(run['forecast_order_inputs_id'],orders['id'])
                saved_inputs=view['views']['/datasets/'+run['dataset_id']]
                self.assertEqual(saved_inputs['sources']['future'],assumptions['id'])
                self.assertEqual(run['input_manifest']['settings']['drivers'],['fx_assumption'])
                self.assertEqual(run['forecast_group_id'],state['group']['id'])
            _,base,_,_=self.command('GET','/runs/'+base_id,None,repeated)
            self.assertEqual(base,{k:v for k,v in original.items() if k!='job_owner'})
            outsider=self.who|{'subject':'another-planner'}
            response,_,_,_=self.command('GET',path,None,repeated,who=outsider)
            self.assertEqual(response.json()['api_status'],404)
        finally:fixture.tearDown()

    def test_monthly_cold_recovery_keeps_factors_and_requires_current_orders_for_export(self):
        from copy import deepcopy
        from datetime import timedelta
        from io import BytesIO
        import pandas as pd
        import shutil
        fixture=sales_fixture.PublicSalesTests();fixture.setUp()
        try:
            source,dataset,_=fixture.history();orders=fixture.orders(dataset)
            parent,_=fixture.calculate(dataset,orders,['model:Last observed'])
            base_id=parent['jobs'][0]['run_id'];original=deepcopy(fixture.ws.load_run(base_id))
            original_bytes=fixture.ws.datasets.source(source['id'])[1]
            history=pd.read_csv(BytesIO(original_bytes),dtype={'sku':str})
            history['fx_assumption']=history.groupby('customer').cumcount()+100
            history.loc[(history.customer=='Aftab')&(history.date==history.date.max()),'quantity']=26
            uploaded=fixture.ws.datasets.upload('corrected-factor-history.csv',history.to_csv(index=False).encode(),'history')
            future=pd.DataFrame({'date':pd.date_range(fixture.today.replace(day=1),periods=3,freq='MS').strftime('%Y-%m-%d'),
                                 'fx_assumption':[136,138,140]})
            future_source=fixture.ws.datasets.upload('reviewed-factor-assumptions.csv',future.to_csv(index=False).encode(),'future')
            settings=dataset['settings']|{'drivers':['fx_assumption'],'future_date_col':'date',
                'future_driver_policy':'require','evidence_policy':'reviewed_what_if'}
            corrected=fixture.ws.datasets.save('Tehran corrected history and factors',
                {'history':uploaded['id'],'future':future_source['id']},settings,'synthetic_sample',True,parent_dataset_id=dataset['id'])
            shutil.copytree(fixture.root/'companies',self.root/'monthly-input/data/companies')
            archive=self.root/'monthly-input.zip';capture(self.root/'monthly-input','tehran_a',archive)
            checkpoint=archive.read_bytes()
            response,state,_,checkpoint=self.command('POST','/forecast-updates',
                {'run_id':base_id,'request_id':'monthly-cold-start'},checkpoint)
            self.assertEqual(response.json()['api_status'],201,state)
            path='/forecast-updates/'+state['id'];sequence=0
            def step(action,**values):
                nonlocal checkpoint,state,sequence
                sequence+=1
                body={'revision':state['revision'],'request_id':'monthly-step-'+str(sequence), 'action':action,**values}
                response,new,_,updated=self.command('POST',path+'/steps',body,checkpoint)
                self.assertEqual(response.json()['api_status'],200,new)
                state,checkpoint=new,updated
                # Every read restores a fresh company from the durable checkpoint.
                response,resumed,_,unchanged=self.command('GET',path,None,checkpoint)
                self.assertEqual(response.json()['api_status'],200,resumed)
                self.assertEqual(resumed['revision'],state['revision']);self.assertEqual(unchanged,checkpoint)
                return body
            body=step('history',dataset_id=corrected['id'],reviewed=True)
            response,_,_,_=self.command('POST',path+'/steps',body|{'request_id':'monthly-stale-tab'},checkpoint)
            self.assertEqual(response.json()['api_status'],400,'Another tab cannot replay an old revision')
            response,_,_,_=self.command('POST',path+'/steps',{'action':'calculate','revision':state['revision'],
                'request_id':'monthly-missing-future','months':4,'method':'model:Ridge + drivers','reviewed':True},checkpoint)
            self.assertEqual(response.json()['api_status'],400,'No fabricated fourth-month factor value')
            step('calculate',months=3,method='model:Ridge + drivers',reviewed=True)
            self.assertEqual(state['job']['state'],'succeeded',state)
            step('calculated');new_id=state['run_id']
            step('no_factors',reviewed=True)
            self.assertEqual(state['stage'],'orders')
            response,_,_,_=self.command('GET',path+'/export?mode=combined_demand&kind=json',None,checkpoint)
            self.assertEqual(response.json()['api_status'],400,'Opening a result does not approve orders')
            inputs=deepcopy(fixture.ws.sales.list(base_id)[0]['inputs']);inputs['run_id']=new_id
            inputs['valid_until']=str(fixture.today+timedelta(days=7))
            def save_orders(request_id):
                nonlocal checkpoint
                body={'inputs':inputs,'request_id':request_id};target='/runs/'+new_id+'/order-snapshots'
                response,review,_,checkpoint=self.command('POST',target+'/preview',body,checkpoint)
                self.assertEqual(response.json()['api_status'],200,review)
                response,saved,_,checkpoint=self.command('POST',target,body|{'review_token':review['review_token']},checkpoint)
                self.assertEqual(response.json()['api_status'],201,saved)
                return saved
            saved=save_orders('monthly-orders-first');step('orders',snapshot_id=saved['id'])
            report=state['comparison'];self.assertEqual(report['overlapping_rows'],12)
            step('review',reviewed=True,review_token=report['review_token']);self.assertTrue(state['ready'])
            expected={(r['customer'],r['sku'],r['period']):r['after_to_serve'] for r in report['rows']}
            for kind in ('json','csv','xlsx'):
                response,_,_,unchanged=self.command('GET',path+'/export?mode=combined_demand&kind='+kind,None,checkpoint)
                self.assertEqual(response.json()['api_status'],200);self.assertEqual(unchanged,checkpoint)
                raw=self.client.get('/output/'+self.attempts[-1]+'/artifact/api-download.bin').content
                if kind=='json':rows=json.loads(raw)
                elif kind=='csv':rows=pd.read_csv(BytesIO(raw),dtype={'sku':str}).to_dict('records')
                else:rows=pd.read_excel(BytesIO(raw),dtype={'sku':str}).to_dict('records')
                self.assertEqual(len(rows),12)
                for row in rows:self.assertAlmostEqual(row['quantity'],expected[(row['customer'],row['sku'],row['period'])],places=9)
            # A later order review invalidates the old approval even after cold restore.
            inputs['orders'][0]['ordered']=30
            newer=save_orders('monthly-orders-newer')
            response,resumed,_,_=self.command('GET',path,None,checkpoint)
            self.assertFalse(resumed['ready']);self.assertIn('Orders changed',resumed['attention'])
            response,_,_,_=self.command('GET',path+'/export?mode=combined_demand&kind=json',None,checkpoint)
            self.assertEqual(response.json()['api_status'],400)
            step('orders',snapshot_id=newer['id']);step('review',reviewed=True,review_token=state['comparison']['review_token'])
            with patch('app.sales_demand.run_today',return_value=fixture.today+timedelta(days=8)):
                response,resumed,_,_=self.command('GET',path,None,checkpoint)
                self.assertFalse(resumed['ready']);self.assertIn('expired',resumed['attention'])
            response,_,_,_=self.command('GET',path,None,checkpoint,who=dict(self.who,subject='another-planner'))
            self.assertEqual(response.json()['api_status'],404)
            restored=self.root/'monthly-restored.zip';restored.write_bytes(checkpoint)
            restore(restored,self.root/'monthly-restored','tehran_a')
            workspaces=CompanyWorkspaces(self.root/'monthly-restored/data/companies')
            try:
                ws=workspaces.for_principal(self.who);run=ws.load_run(new_id)
                self.assertEqual(run['input_manifest']['settings']['drivers'],['fx_assumption'])
                self.assertEqual(run['input_manifest']['settings']['evidence_policy'],'reviewed_what_if')
                self.assertEqual(ws.datasets.get(run['dataset_id'])['sources']['future'],future_source['id'])
                self.assertEqual(ws.datasets.source(future_source['id'])[1],future.to_csv(index=False).encode())
                self.assertEqual(ws.load_run(base_id),original)
                self.assertEqual(ws.datasets.source(source['id'])[1],original_bytes)
            finally:workspaces.close()
        finally:fixture.tearDown()

    def test_monthly_stopped_job_recovers_after_restore_without_replaying_calculation(self):
        from app.company_context import personal_owner
        from app.company_workflows import monthly
        from app.monthly_refresh import StartUpdate,UpdateStep
        import shutil
        fixture=sales_fixture.PublicSalesTests();fixture.setUp()
        try:
            _,dataset,_=fixture.history();group,_=fixture.calculate(dataset,fixture.orders(dataset),['model:Last observed'])
            service=monthly(fixture.ws,lambda *_:None);owner=personal_owner(self.who)
            state=service.start(owner,StartUpdate(run_id=group['jobs'][0]['run_id'],request_id='interrupted-monthly-start'))
            state=service.step(state['id'],owner,UpdateStep(revision=0,action='history',dataset_id=dataset['id'],
                reviewed=True,request_id='interrupted-monthly-history'))
            state=service.step(state['id'],owner,UpdateStep(revision=state['revision'],action='calculate',months=2,
                method='model:Recent average',reviewed=True,request_id='interrupted-monthly-calculate'))
            fixture.ws.jobs.cancel(state['job_id'])
            shutil.copytree(fixture.root/'companies',self.root/'stopped-input/data/companies')
            archive=self.root/'stopped-input.zip';capture(self.root/'stopped-input','tehran_a',archive);checkpoint=archive.read_bytes()
            path='/forecast-updates/'+state['id']
            response,resumed,_,_=self.command('GET',path,None,checkpoint)
            self.assertEqual(response.json()['api_status'],200,resumed);self.assertEqual(resumed['job']['state'],'cancelled')
            response,state,_,checkpoint=self.command('POST',path+'/steps',{'revision':state['revision'],
                'action':'retry_calculation','request_id':'restore-stopped-review'},checkpoint)
            self.assertEqual(response.json()['api_status'],200,state);self.assertEqual(state['stage'],'forecast')
            self.assertEqual((state['months'],state['method']),(2,'model:Recent average'))
            body={'revision':state['revision'],'action':'calculate','months':2,'method':'model:Recent average',
                  'request_id':'restore-calculation-retry','reviewed':False}
            response,_,_,_=self.command('POST',path+'/steps',body,checkpoint)
            self.assertEqual(response.json()['api_status'],400,'A stopped job cannot bypass a fresh confirmation')
            body['reviewed']=True
            response,state,view,checkpoint=self.command('POST',path+'/steps',body,checkpoint)
            self.assertEqual(response.json()['api_status'],200,state);self.assertEqual(state['job']['state'],'succeeded')
            before=len(view['views']['/jobs']['jobs'])
            response,retried,view,checkpoint=self.command('POST',path+'/steps',body,checkpoint)
            self.assertEqual(response.json()['api_status'],200,retried)
            self.assertEqual(retried['job']['id'],state['job']['id']);self.assertEqual(len(view['views']['/jobs']['jobs']),before)
        finally:fixture.tearDown()

    def test_settings_and_personal_views_use_original_permission_and_owner_gates(self):
        admin=dict(self.who,role='admin',permissions=self.who['permissions']+['settings:manage'])
        response,site,view,checkpoint=self.command('PUT','/settings/site',{'name':'Tehran test factory','province':'Tehran','timezone':'Asia/Tehran'},who=admin)
        self.assertEqual(response.json()['api_status'],200,site)
        self.assertEqual(view['views']['/workspace']['site']['name'],'Tehran test factory')
        response,_,_,_=self.command('PUT','/settings/site',{'name':'Wrong','province':'Tehran','timezone':'Asia/Tehran'},checkpoint)
        self.assertEqual(response.status_code,422)
        key=dict(self.who,subject='service:test-key',auth_kind='api_key',key_kind='company',permissions=self.who['permissions']+['views:own'])
        response,value,_,_=self.command('GET','/views?run_id=unknown',None,checkpoint,who=key)
        self.assertIn(response.json()['api_status'],[403,404])

    def test_saved_conversations_remain_personal_and_survive_rename_archive_restore(self):
        from app.company_context import personal_owner
        who=dict(self.who,permissions=self.who['permissions']+['ai:query','chats:own'])
        workspace=CompanyWorkspaces(self.root/'chat/data/companies')
        try:
            w=workspace.for_principal(who)
            chat=w.journal.put(personal_owner(who),{'question':'Check Tehran demand','answer':'Review saved demand.',
                'run_id':None,'snapshot_id':None,'dataset_id':None,'actions':[]})
            archive=self.root/'chat.zip';capture(self.root/'chat','tehran_a',archive)
        finally:workspace.close()
        checkpoint=archive.read_bytes()
        response,_,view,checkpoint=self.command('POST','/ai/conversations/'+chat+'/rename',{'title':'Tehran sales review'},checkpoint,who=who)
        self.assertEqual(response.json()['api_status'],200)
        self.assertFalse(any(path.startswith('/ai') for path in view['views']),'Private chats never enter shared projections')
        response,listed,_,_=self.command('GET','/ai/conversations',None,checkpoint,who=who)
        self.assertEqual(listed['chats'][0]['title'],'Tehran sales review')
        other=dict(who,subject='other-user')
        response,listed,_,_=self.command('GET','/ai/conversations',None,checkpoint,who=other)
        self.assertEqual(listed['chats'],[])
        response,_,_,_=self.command('GET','/ai/conversations/'+chat,None,checkpoint,who=other)
        self.assertEqual(response.json()['api_status'],400)
        response,_,_,checkpoint=self.command('POST','/ai/conversations/'+chat+'/manage',{'operation':'archive'},checkpoint,who=who)
        self.assertEqual(response.json()['api_status'],200)
        response,listed,_,_=self.command('GET','/ai/conversations?status=archived',None,checkpoint,who=who)
        self.assertEqual(len(listed['chats']),1)
        response,_,_,checkpoint=self.command('POST','/ai/conversations/'+chat+'/manage',{'operation':'restore'},checkpoint,who=who)
        self.assertEqual(response.json()['api_status'],200)
        response,descriptor,_,_=self.command('GET','/ai/conversations/'+chat+'/export',None,checkpoint,who=who)
        self.assertEqual(response.json()['api_status'],200,descriptor)
        self.assertIn(b'Review saved demand',self.client.get('/output/'+self.attempts[-1]+'/artifact/api-download.bin').content)
        service=dict(who,subject='service:integration',auth_kind='api_key',key_kind='company')
        response,_,_,_=self.command('GET','/ai/conversations',None,checkpoint,who=service)
        self.assertEqual(response.json()['api_status'],403)
