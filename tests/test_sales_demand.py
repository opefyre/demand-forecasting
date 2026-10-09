from copy import deepcopy
from datetime import date, timedelta
from pathlib import Path
from io import BytesIO
import json
import tempfile
import unittest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from app.sales_demand import DemandStore, demand_outlook, export_demand, import_rows
from app.sales_api import install_sales_routes
from app.datasets import DatasetStore


TODAY = date.today()
PERIOD = str(TODAY.replace(day=1))


def fixture():
    run = {'run_id': 'forecast1', 'source_classification': 'synthetic_sample', 'unit': 'tonnes',
           'run_settings': {'frequency': 'monthly'}, 'metadata': {}, 'series': {}}
    customers = []
    for customer, quantity in [('A', 10), ('B', 8), ('C', 5)]:
        run['metadata'][customer] = {'customer': customer, 'sku': '001'}
        run['series'][customer] = {'forecast': [{'timestamp': PERIOD, 'mean': quantity}]}
        customers.append({'customer': customer, 'sku': '001', 'unit': 'tonnes', 'series_id': customer})
    inputs = {'name':'Reviewed orders', 'run_id':'forecast1', 'classification':'synthetic_sample',
              'as_of':str(TODAY), 'valid_until':str(TODAY + timedelta(days=7)),
              'order_feed':'complete_snapshot', 'customers':customers, 'orders':[],
              'commitments':[], 'reviewed':True, 'note':'Checked full customer list.'}
    return run, inputs


def order(customer='A', quantity=6, **kw):
    return {'reference':customer+'/1', 'customer':customer, 'sku':'001', 'unit':'tonnes',
            'due_date':str(TODAY), 'ordered':quantity, 'fulfilled':0, 'cancelled':0,
            'status':'confirmed', **kw}


class DemandMathTests(unittest.TestCase):
    def setUp(self): self.run, self.inputs = fixture()
    def outlook(self): return demand_outlook(self.inputs, self.run)
    def first(self): return self.outlook()['rows'][0]

    def test_orders_above_forecast_not_added_twice(self):
        self.inputs['orders'] = [order(quantity=16)]
        r = self.first()
        self.assertEqual((r['baseline'],r['booked'],r['remaining'],r['total']),(10,16,0,16))

    def test_partial_order_keeps_remainder(self):
        self.inputs['orders'] = [order()]
        r = self.first()
        self.assertEqual((r['booked'],r['remaining'],r['total']),(6,4,10))

    def test_full_commitment_can_replace_baseline(self):
        self.inputs['orders'] = [order()]
        self.inputs['commitments'] = [{'customer':'A','sku':'001','unit':'tonnes','period':PERIOD,
            'quantity':6,'owner':'Planner','reason':'Customer confirms entire monthly demand',
            'valid_until':str(TODAY+timedelta(days=7))}]
        r = self.first()
        self.assertEqual((r['baseline'],r['remaining'],r['total']),(10,0,6))
        self.inputs['commitments'][0]['quantity'] = 5
        self.assertIsNone(self.first()['total'])
        self.assertFalse(self.outlook()['can_export'])

    def test_customer_without_orders_is_retained(self):
        self.inputs['orders'] = [order(quantity=16),order('C',3)]
        r = self.outlook()['rows']
        self.assertEqual(sum(x['booked'] for x in r),19)
        self.assertEqual(sum(x['remaining'] for x in r),10)
        self.assertEqual(sum(x['total'] for x in r),29)
        self.assertEqual(r[1]['total'],8)

    def test_current_month_fulfillment(self):
        self.inputs['orders'] = [order(quantity=8,fulfilled=3)]
        r = self.first()
        self.assertEqual((r['fulfilled'],r['booked'],r['remaining'],r['still_to_serve']),(3,5,2,7))

    def test_no_history_is_unknown_not_zero(self):
        self.inputs['customers'].append({'customer':'New','sku':'001','unit':'tonnes'})
        self.inputs['orders'] = [order('New',16)]
        r = self.outlook()['rows'][-1]
        self.assertEqual(r['booked'],16)
        self.assertIsNone(r['remaining'])
        self.assertFalse(self.outlook()['can_export'])

    def test_unknown_or_expired_order_source_blocks_export(self):
        self.inputs['order_feed'] = 'unknown'
        self.assertIsNone(self.first()['total'])
        self.inputs['order_feed'] = 'complete_snapshot'
        later = demand_outlook(self.inputs,self.run,today=TODAY+timedelta(days=8))
        self.assertFalse(later['can_export'])
        with self.assertRaises(ValueError): export_demand(later,'remaining_forecast','csv')

    def test_cancellation_and_split_schedules(self):
        self.inputs['orders'] = [order(quantity=12,cancelled=4),order(quantity=4,reference='A/2')]
        self.assertEqual(self.first()['total'],12)
        self.inputs['orders'][0].update(status='cancelled',cancelled=9,fulfilled=3)
        self.assertEqual((self.first()['fulfilled'],self.first()['booked']),(3,4))

    def test_invalid_status_quantities_and_duplicates_rejected(self):
        for records in ([order(),order()], [order(quantity=4,fulfilled=5)],
                        [order(status='cancelled')], [order(status='unconfirmed',fulfilled=1)]):
            self.inputs['orders'] = records
            with self.assertRaises(ValueError): self.outlook()

    def test_scope_mismatches_rejected(self):
        for changes in ({'customer':'Other'}, {'unit':'kg'}, {'sku':'002'}):
            self.inputs['orders'] = [order(**changes)]
            with self.assertRaises(ValueError): self.outlook()
        self.inputs['orders']=[]
        self.inputs['customers'][0]['series_id']='B'
        with self.assertRaises(ValueError): self.outlook()

    def test_unmapped_series_and_backlog_block_export(self):
        self.inputs['customers'].pop()
        self.assertFalse(self.outlook()['can_export'])
        self.run,self.inputs=fixture()
        self.inputs['orders']=[order(due_date=str(TODAY.replace(day=1)-timedelta(days=1)))]
        self.assertFalse(self.outlook()['can_export'])
        self.assertTrue(any('Past-due' in w for w in self.outlook()['warnings']))

    def test_exports_residual_and_combined_exclude_fulfilled(self):
        self.inputs['orders']=[order(quantity=8,fulfilled=3)]
        output=self.outlook()
        residual=json.loads(export_demand(output,'remaining_forecast','json')[0])
        combined=json.loads(export_demand(output,'combined_demand','json')[0])
        self.assertEqual(residual[0]['quantity'],2)
        self.assertEqual(combined[0]['quantity'],7)
        self.assertEqual(combined[0]['approval'],'draft')
        import openpyxl
        book=openpyxl.load_workbook(BytesIO(export_demand(output,'combined_demand','xlsx')[0]))
        self.assertEqual(book.active.max_row,4)
        book.close()


class DemandPersistenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.path=Path(self.tmp.name)
        self.sources=DatasetStore(self.path/'sources');self.store=DemandStore(self.path/'sales.sqlite3')
        self.run,self.inputs=fixture()
        app=FastAPI()
        @app.middleware('http')
        async def principal(request:Request,call_next):
            request.state.principal={'issuer':'test','subject':'alice'}
            return await call_next(request)
        install_sales_routes(app,self.store,self.sources,lambda _: self.run)
        self.client=TestClient(app)
    def tearDown(self): self.tmp.cleanup()

    def test_import_preserves_identifiers_and_cells(self):
        source=self.sources.upload('orders.csv',b'reference,customer,sku,unit,due_date,ordered,status\n0001,A,001,tonnes,'+str(TODAY).encode()+b',6,confirmed\n','sales_orders')
        cfg={'source_id':source['id'],'mapping':dict(zip(['reference','customer','sku','unit','due_date','ordered','status'],'ABCDEFG'))}
        rows,evidence=import_rows(self.sources,'orders',cfg)
        self.assertEqual((rows[0]['reference'],rows[0]['sku']),('0001','001'))
        self.assertEqual(evidence['cells'][0]['cells']['ordered'],'F2')
        cfg['mapping']['ordered']='A'
        with self.assertRaises(ValueError): import_rows(self.sources,'orders',cfg)

    def test_api_idempotency_and_baseline_fingerprint(self):
        payload={'inputs':self.inputs,'request_id':'sales-save-001'}
        response=self.client.post('/api/sales/inputs',json=payload)
        self.assertEqual(response.status_code,200,response.text)
        saved=response.json()
        self.assertIn('alice',saved['actor'])
        self.assertEqual(self.client.post('/api/sales/inputs',json=payload).json()['id'],saved['id'])
        self.assertEqual(len(self.store.list('forecast1')),1)
        self.assertEqual(self.client.get(f'/api/sales/inputs/{saved["id"]}/outlook').status_code,200)
        payload['inputs']['note']='Changed review note'
        self.assertEqual(self.client.post('/api/sales/inputs',json=payload).status_code,400)
        self.run['series']['A']['forecast'][0]['mean']=999
        self.assertEqual(self.client.get(f'/api/sales/inputs/{saved["id"]}/outlook').status_code,400)

    def test_source_role_and_sample_real_separation(self):
        self.assertEqual(self.client.post('/api/sales/sources',data={'role':'orders'},files={'file':('orders.csv',b'reference,customer\n')}).status_code,200)
        self.run['source_classification']='user_provided'
        self.assertEqual(self.client.get('/api/sales/runs/forecast1/sample').status_code,400)
        self.assertEqual(self.client.post('/api/sales/validate',json={'inputs':self.inputs}).status_code,400)

    def test_formulas_rejected(self):
        import openpyxl
        book=openpyxl.Workbook();sheet=book.active
        sheet.append(['customer','sku','unit','series_id']);sheet.append(['A','001','tonnes','=1+1'])
        blob=BytesIO();book.save(blob);book.close()
        source=self.sources.upload('customers.xlsx',blob.getvalue(),'sales_customers')
        with self.assertRaisesRegex(ValueError,'formulas'):
            import_rows(self.sources,'customers',{'source_id':source['id'],'mapping':dict(zip(['customer','sku','unit','series_id'],'ABCD'))})
