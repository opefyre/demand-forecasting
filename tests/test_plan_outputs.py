from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from fastapi.testclient import TestClient

import app.main as main
from app.datasets import DatasetStore
from app.planning import PlanStore
from app.plan_outputs import resolve_plan, plan_workbook


class PlanOutputTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = PlanStore(Path(self.tmp.name) / 'plans.json')
        self.data = DatasetStore(Path(self.tmp.name) / 'inputs')
        self.source = self.data.upload('production.xlsx', (main.SAMPLES_DIR / 'iran_operations_master.xlsx').read_bytes(), 'operations')
        self.run = {'run_id': 'test', 'unit': 'tonnes', 'run_settings': {'frequency': 'monthly'},
                    'metadata': {'A': {'sku': 'PKG-KRAFT-120', 'production_line': 'Paper line A'}},
                    'series': {'A': {'forecast': [{'timestamp': '2026-09-01', 'mean': 250., 'p10': 200., 'p90': 300.}]},
                               '__all__': {'forecast': [{'timestamp': '2026-09-01', 'mean': 250.}]}},
                    'input_manifest': {'sources': [{'role': 'operations', 'id': self.source['id'], 'sha256': self.source['sha256']}]}}
        self.plan = self.store.create(name='Test', run_id='test', site_id='site', owner='Planner', settings={}, metrics={'evidence_level': 'strong'})
        self.plan = self.store.add_override(self.plan['id'], item_id='A', period='2026-09-01', value=500, reason='Confirmed orders', actor='Planner')

    def tearDown(self):
        self.tmp.cleanup()

    def test_resolves_quantities_without_changing_baseline(self):
        original = deepcopy(self.run)
        resolved, rows = resolve_plan(self.run, self.plan)
        self.assertEqual(rows[0]['plan_quantity'], 500)
        self.assertEqual(rows[0]['forecast_quantity'], 250)
        self.assertEqual(resolved['series']['__all__']['forecast'][0]['mean'], 500)
        self.assertNotIn('p90', resolved['series']['A']['forecast'][0])
        self.assertEqual(self.run, original)
        self.plan['overrides'][0]['reverted_at'] = '2026-09-21'
        self.assertEqual(resolve_plan(self.run, self.plan)[1][0]['plan_quantity'], 250)

    def test_latest_active_adjustment_wins_and_bad_quantities_fail(self):
        self.plan['overrides'].append({**self.plan['overrides'][0], 'value': 600})
        self.assertEqual(resolve_plan(self.run, self.plan)[1][0]['plan_quantity'], 600)
        for value in (float('nan'), float('inf'), -1):
            with self.assertRaises(ValueError):
                self.store.add_override(self.plan['id'], item_id='A', period='2026-09-01', value=value, reason='test', actor='test')

    def test_export_contains_exact_plan_and_safe_text(self):
        self.plan['overrides'][0]['reason'] = '=HYPERLINK("http://invalid")'
        _, rows = resolve_plan(self.run, self.plan)
        book = openpyxl.load_workbook(BytesIO(plan_workbook(self.plan, rows)))
        sheet = book['Plan quantities']
        values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
        self.assertEqual(values['plan_quantity'], 500)
        self.assertEqual(values['forecast_quantity'], 250)
        self.assertFalse(any(c.data_type == 'f' for s in book for row in s for c in row))

    def test_supply_and_export_share_approved_quantities_and_hash_guard(self):
        with patch.object(main, 'PLAN_STORE', self.store), patch.object(main, 'DATASET_STORE', self.data), patch.object(main, '_load_run', return_value=self.run), TestClient(main.app) as client:
            url = '/api/plans/' + self.plan['id']
            self.assertEqual(client.get(url + '/supply').status_code, 400)
            self.store.transition(self.plan['id'], status='review', actor='Planner')
            self.store.transition(self.plan['id'], status='approved', actor='Reviewer')
            supply = client.get(url + '/supply')
            self.assertEqual(supply.status_code, 200, supply.text)
            self.assertEqual(supply.json()['capacity'][0]['forecast_tonnes'], 500)
            exported = client.get(url + '/export?include_supply=true')
            self.assertEqual(exported.status_code, 200)
            book = openpyxl.load_workbook(BytesIO(exported.content))
            values = list(book['Capacity'].values)
            self.assertEqual(dict(zip(values[0], values[1]))['forecast_tonnes'], 500)
            self.run['input_manifest']['sources'][0]['sha256'] = 'changed'
            self.assertEqual(client.get(url + '/supply').status_code, 400)
            self.assertEqual(client.get(url + '/export?include_supply=true').status_code, 400)

    def test_review_quantities_match_export_and_reversal(self):
        with patch.object(main, 'PLAN_STORE', self.store), patch.object(main, '_load_run', return_value=self.run), TestClient(main.app) as client:
            url = '/api/plans/' + self.plan['id']
            review = client.get(url + '/quantities')
            self.assertEqual(review.status_code, 200, review.text)
            payload = review.json()
            self.assertEqual(payload['plan_id'], self.plan['id'])
            self.assertEqual(payload['run_id'], 'test')
            self.assertEqual(payload['status'], 'draft')
            self.assertEqual(payload['rows'][0]['plan_quantity'], 500)
            self.assertEqual(payload['rows'][0]['forecast_quantity'], 250)
            self.assertEqual(payload['rows'][0]['adjustment'], 250)
            book = openpyxl.load_workbook(BytesIO(client.get(url + '/export').content))
            sheet = book['Plan quantities']
            exported = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
            self.assertEqual(exported['plan_quantity'], payload['rows'][0]['plan_quantity'])
            book.close()
            override = self.plan['overrides'][0]['id']
            response = client.post(url + '/overrides/' + override + '/revert', json={'reason': 'Order withdrawn', 'actor': 'Planner'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(client.get(url + '/quantities').json()['rows'][0]['plan_quantity'], 250)
            self.assertEqual(self.run['series']['A']['forecast'][0]['mean'], 250)

    def test_review_rejects_missing_plan_or_invalid_override(self):
        with patch.object(main, 'PLAN_STORE', self.store), patch.object(main, '_load_run', return_value=self.run), TestClient(main.app) as client:
            self.assertEqual(client.get('/api/plans/missing/quantities').status_code, 404)
            altered = deepcopy(self.plan)
            altered['overrides'][0]['period'] = '2030-01-01'
            with patch.object(self.store, 'get', return_value=altered):
                self.assertEqual(client.get('/api/plans/' + self.plan['id'] + '/quantities').status_code, 400)


if __name__ == '__main__':
    unittest.main()
