import asyncio
from copy import deepcopy
import csv
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import unittest
import uuid
from unittest.mock import patch

import pandas as pd
import openpyxl
from fastapi.testclient import TestClient

from tests import test_factor_links as fixtures
from app.factor_links import preview_link,save_link
from app.factor_batch import preview_batch,save_batch,check_saved_batch,compose_batch,calculate_batch
from app.sales_demand import demand_outlook,export_demand


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.FactorLinkTests();self.f.setUp();self.addCleanup(self.f.tearDown)
        self.base=deepcopy(self.f.base)
        self.base.update(unit='tonnes',metadata={'A':{'customer':'A','sku':'SKU'},'B':{'customer':'B','sku':'SKU'}})
        self.a=self.group('A',150);self.b=self.group('B',250)

    def group(self,series,amount):
        body={**self.f.payload,'future_value':amount,'series_ids':[series]}
        report=preview_link(self.base,self.f.store,self.f.factors,body)
        return save_link(self.base,self.f.store,self.f.factors,{**body,'reviewed':True,
            'review_token':report['review_token'],'request_id':str(uuid.uuid4())})

    def report(self,ids=None):
        return preview_batch(self.base,self.f.store,self.f.factors,{'dataset_ids':[self.a['id'],self.b['id']] if ids is None else ids})

    def save(self,report=None,ids=None,request=None):
        return save_batch(self.base,self.f.store,self.f.factors,{
            'dataset_ids':ids or [self.a['id'],self.b['id']],'reviewed':True,
            'review_token':(report or self.report(ids))['review_token'],'request_id':request or str(uuid.uuid4())})

    def test_distinct_assumptions_scopes_and_idempotent_immutable_save(self):
        original=deepcopy(self.f.dataset);before=deepcopy(self.base);token=self.report();key=str(uuid.uuid4())
        saved=self.save(token,request=key);self.assertEqual(saved,self.save(token,request=key))
        self.assertEqual(token['series_count'],2);self.assertEqual(token['unchanged_series_ids'],[])
        self.assertEqual([g['alignment']['future_value'] for g in token['groups']],[150,250])
        self.assertFalse(saved['scenario_provenance']['orders_changed'])
        self.assertEqual(self.f.store.get(original['id']),original);self.assertEqual(self.base,before)
        self.assertEqual(check_saved_batch(self.base,saved,self.f.store,self.f.factors),token)
        with self.assertRaises(ValueError):self.save(ids=[self.a['id']],request=key)

    def test_baseline_retained_and_overlap_duplicates_and_strict_payload_rejected(self):
        self.assertEqual(self.report([self.a['id']])['unchanged_series_ids'],['B'])
        other=self.group('A',200)
        for ids in ([self.a['id'],other['id']],[self.a['id'],self.a['id']],[],['unknown']):
            with self.subTest(ids=ids),self.assertRaises(ValueError):self.report(ids)
        for payload in ({'dataset_ids':[True]},{'dataset_ids':[self.a['id']],'method':'recommended'}):
            with self.assertRaises(ValueError):preview_batch(self.base,self.f.store,self.f.factors,payload)

    def test_foreign_unscoped_older_and_unreviewed_groups_block(self):
        original_get=self.f.store.get
        variants=[]
        for key,value in [('base_run_id','foreign'),('type','factor_batch'),('reviewed_inputs',None)]:
            changed=deepcopy(self.a);changed['scenario_provenance'][key]=value;variants.append(changed)
        changed=deepcopy(self.a);changed['scenario_provenance']['alignment'].pop('series_ids');variants.append(changed)
        for changed in variants:
            with self.subTest(value=changed),patch.object(self.f.store,'get',side_effect=lambda i:changed if i==changed['id'] else original_get(i)),self.assertRaises(ValueError):
                self.report([self.a['id']])
        report=self.report()
        with self.assertRaises(ValueError):save_batch(self.base,self.f.store,self.f.factors,{
            'dataset_ids':[self.a['id'],self.b['id']],'reviewed':False,
            'review_token':report['review_token'],'request_id':str(uuid.uuid4())})

    def test_changed_baseline_snapshot_and_assumptions_require_new_review(self):
        report=self.report();self.base['metadata']['A']['customer']='Changed'
        with self.assertRaises(ValueError):self.save(report)
        self.base['metadata']['A']['customer']='A'
        (self.f.store.root/(self.f.snapshot['source']['id']+'.bin')).write_bytes(b'changed')
        with self.assertRaises(ValueError):self.report()

    def test_partial_batch_failure_publishes_no_combined_result(self):
        saved=self.save();calls=[]
        async def fail(key):
            calls.append(key)
            if len(calls)==2:raise ValueError('Synthetic failure')
            return {'synthetic_completed_first_group':True}
        folder=self.f.root/'batch';folder.mkdir()
        with self.assertRaisesRegex(ValueError,'Synthetic failure'):
            asyncio.run(calculate_batch(self.base,saved,self.f.store,self.f.factors,None,None,fail,folder))
        self.assertEqual(calls,[self.a['id'],self.b['id']]);self.assertEqual(list(folder.iterdir()),[])

    def test_changed_derived_settings_and_late_source_change_block(self):
        changed=deepcopy(self.a);changed['settings']['horizon']=12;get=self.f.store.get
        with patch.object(self.f.store,'get',side_effect=lambda i:changed if i==changed['id'] else get(i)),self.assertRaisesRegex(ValueError,'inputs changed'):
            self.report()
        saved=self.save();folder=self.f.root/'late-change';folder.mkdir()
        async def change(key):
            (self.f.store.root/(self.f.snapshot['source']['id']+'.bin')).write_bytes(b'changed')
            return {}
        with self.assertRaises(ValueError):asyncio.run(calculate_batch(self.base,saved,self.f.store,self.f.factors,None,None,change,folder))
        self.assertEqual(list(folder.iterdir()),[])

    def test_actual_api_validates_review_and_changed_batch(self):
        from app import main
        with patch.object(main,'_load_run',return_value=self.base),patch.object(main,'DATASET_STORE',self.f.store),patch.object(main,'FACTOR_STORE',self.f.factors),TestClient(main.app) as client:
            path='/api/runs/base/factor-batch'
            response=client.post(path+'/preview',json={'dataset_ids':[self.a['id']]})
            self.assertEqual(response.status_code,200,response.text)
            body={'dataset_ids':[self.a['id']],'reviewed':True,'review_token':response.json()['review_token'],'request_id':str(uuid.uuid4())}
            self.assertEqual(client.post(path,json=body).status_code,200)
            self.assertEqual(client.post(path,json={**body,'dataset_ids':[self.b['id']]}).status_code,400)
            self.assertEqual(client.post(path+'/preview',json={'dataset_ids':[True]}).status_code,400)

    def test_real_engine_two_groups_untouched_product_orders_and_six_exports(self):
        from app import main
        c=self.f.original[self.f.original.item.eq('B')].copy();c['item']='C';c['customer']='C';c['qty']=80
        history=pd.concat([self.f.original,c],ignore_index=True)
        source=self.f.store.upload('three-customers.csv',history.to_csv(index=False).encode(),'history')
        original=self.f.store.save('Three synthetic customers',{'history':source['id']},self.f.dataset['settings'],'synthetic_sample',True)
        runs=self.f.root/'runs';runs.mkdir()
        with patch.object(main,'DATASET_STORE',self.f.store),patch.object(main,'FACTOR_STORE',self.f.factors),patch.object(main,'RUNS_DIR',runs):
            self.base=asyncio.run(main.run_saved(main.SavedRunConfig(dataset_id=original['id'])))
            untouched=deepcopy(self.base);self.a=self.group('A',150);self.b=self.group('B',250)
            saved=self.save();combined=asyncio.run(main.run_saved(main.SavedRunConfig(dataset_id=saved['id'])))
            self.assertEqual(combined['series']['C'],self.base['series']['C'])
            self.assertEqual(main._load_run(self.base['run_id']),untouched)
            self.assertEqual(combined['scenario']['type'],'factor_batch');self.assertEqual(len(combined['forecast_rows']),6)
            self.assertIsNone(combined['metrics']['wape_pct']);self.assertEqual(combined['range_model'],{})
            self.assertFalse(combined['metrics']['automatic_publish_allowed'])
            for point in combined['series']['__all__']['forecast']:
                self.assertAlmostEqual(point['mean'],sum(next(r['mean'] for r in combined['series'][k]['forecast'] if r['timestamp']==point['timestamp']) for k in ('A','B','C')))
                self.assertIsNone(point['p10'])
            # Every selected forecast equals the separately calculated group, not a second refit of the merged rows.
            children=[]
            for d,customer in ((self.a,'A'),(self.b,'B')):
                separate=asyncio.run(main.run_saved(main.SavedRunConfig(dataset_id=d['id'])))
                children.append(separate)
                self.assertEqual(combined['series'][customer]['forecast'],separate['series'][customer]['forecast'])
            report=self.report()
            for change in ('nonfinite','period','table'):
                bad=deepcopy(children)
                if change=='nonfinite':bad[0]['series']['A']['forecast'][0]['mean']=float('nan')
                elif change=='period':bad[0]['series']['A']['forecast'][0]['timestamp']='2027-01-01'
                else:next(r for r in bad[0]['forecast_rows'] if r['item_id']=='A')['mean']+=1
                with self.subTest(change=change),self.assertRaises(ValueError):compose_batch(self.base,saved,report,bad,runs)
            orders={'name':'Synthetic batch orders','classification':'synthetic_sample','run_id':combined['run_id'],
                'as_of':'2026-01-01','valid_until':'2026-02-01','order_feed':'complete_snapshot',
                'customers':[{'customer':k,'sku':'SKU','unit':'tonnes','series_id':k} for k in ('A','B','C')],
                'orders':[{'reference':'A1','customer':'A','sku':'SKU','unit':'tonnes','due_date':'2026-01-15','ordered':200,'fulfilled':20,'cancelled':0,'status':'confirmed'},
                          {'reference':'B1','customer':'B','sku':'SKU','unit':'tonnes','due_date':'2026-02-15','ordered':10,'fulfilled':0,'cancelled':0,'status':'confirmed'}],
                'commitments':[],'reviewed':True,'note':'Artificial test only.'}
            before=deepcopy(orders);outlook=demand_outlook(orders,combined,today=date(2026,1,1))
            self.assertTrue(outlook['can_export']);self.assertEqual(orders,before)
            expected={}
            for mode in ('combined_demand','remaining_forecast'):
                values=[]
                for kind in ('csv','json','xlsx'):
                    content,_=export_demand(outlook,mode,kind)
                    if kind=='csv':rows=list(csv.DictReader(content.decode('utf-8-sig').splitlines()))
                    elif kind=='json':rows=json.loads(content)
                    else:
                        book=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True)
                        values_sheet=list(book.worksheets[0].values);rows=[dict(zip(values_sheet[0],r)) for r in values_sheet[1:]];book.close()
                    self.assertEqual(len(rows),6)
                    total=sum(float(r['quantity']) for r in rows);values.append(total)
                    independent=sum(max(r['baseline'],r['booked']+r['fulfilled'])-r['fulfilled'] if mode=='combined_demand'
                                    else max(r['baseline']-r['booked']-r['fulfilled'],0) for r in outlook['rows'])
                    self.assertAlmostEqual(total,independent)
                self.assertAlmostEqual(values[0],values[1]);self.assertAlmostEqual(values[1],values[2]);expected[mode]=values[0]
            self.assertAlmostEqual(expected['combined_demand']-expected['remaining_forecast'],190)
            book=openpyxl.load_workbook(runs/combined['run_id']/'forecast_package.xlsx',read_only=True,data_only=True)
            self.assertIn('Customer product factors',book.sheetnames);book.close()
            self.assertEqual(self.f.store.get(original['id']),original)
