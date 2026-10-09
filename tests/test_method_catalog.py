import unittest
from unittest.mock import patch
import warnings
from pathlib import Path
from tempfile import TemporaryDirectory
import json
import openpyxl

import numpy as np
import pandas as pd
from sklearn.exceptions import ConvergenceWarning
from sklearn.linear_model import ElasticNet
from statsforecast.models import AutoETS, Holt, HoltWinters, MSTL

from app.forecast_engine import ModelSpec, _model_specs, _stat_forecast, _make_pipeline, _fit_pipeline, _fit_backtest, _fit_full_models_and_predict, run_forecast


class MethodCatalogTests(unittest.TestCase):
    def test_catalog_contains_promised_library_methods(self):
        for profile in ['fast', 'deep']:
            specs = {s.name: s for s in _model_specs(profile)}
            self.assertTrue({'Holt trend', 'Holt-Winters seasonal', 'Weighted recent average', 'MSTL weekly + yearly', 'Elastic Net + drivers'}.issubset(specs))
            self.assertIsInstance(specs['Elastic Net + drivers'].estimator, ElasticNet)

    def test_weighted_average_is_exact_and_does_not_recurse(self):
        values = _stat_forecast(ModelSpec('Weighted recent average', 'weighted_average'), [999, 10, 20, 30], 8, 12)
        np.testing.assert_allclose(values, [140/6] * 8)
        with self.assertRaisesRegex(ValueError, '3 training periods'):
            _stat_forecast(ModelSpec('Weighted recent average', 'weighted_average'), [10, 20], 8, 12)

    def test_holt_variants_match_library_outputs(self):
        y = np.asarray([80 + i * 0.5 + (i % 12) * 2 for i in range(48)], dtype=float)
        for kind, model in [('holt', Holt(error_type='A')), ('holt_winters', HoltWinters(season_length=12, error_type='A'))]:
            np.testing.assert_allclose(_stat_forecast(ModelSpec(kind, kind), y.tolist(), 12, 12, 'monthly'), np.maximum(0, model.forecast(y=y, h=12)['mean']))
        with self.assertRaisesRegex(ValueError, '24 training periods'):
            _stat_forecast(ModelSpec('Holt-Winters seasonal', 'holt_winters'), [10.] * 23, 6, 12)

    def test_mstl_rejects_wrong_grain_and_insufficient_history(self):
        spec = ModelSpec('MSTL weekly + yearly', 'mstl')
        with self.assertRaisesRegex(ValueError, 'daily observations'):
            _stat_forecast(spec, [10.] * 800, 6, 12, 'monthly')
        with self.assertRaisesRegex(ValueError, '731 training periods'):
            _stat_forecast(spec, [10.] * 730, 14, 7, 'daily')

    def test_mstl_matches_library_and_learns_two_daily_cycles(self):
        t = np.arange(800)
        y = 100 + t*.03 + 10*np.sin(t*2*np.pi/7) + 20*np.sin(t*2*np.pi/365)
        spec = ModelSpec('MSTL weekly + yearly', 'mstl')
        predicted = _stat_forecast(spec, y.tolist(), 14, 7, 'daily')
        expected = MSTL(season_length=[7, 365], trend_forecaster=AutoETS(season_length=1, model='ZZN')).forecast(y=y, h=14)['mean']
        np.testing.assert_allclose(predicted, np.maximum(0, expected))
        self.assertGreater(np.std(predicted), 1)

    def test_mstl_daily_backtest_has_reserved_confirmation(self):
        t = np.arange(850)
        history = pd.DataFrame({'item_id': 'A', 'timestamp': pd.date_range('2023-01-01', periods=850, freq='D'), 'target': 100 + 10*np.sin(t*2*np.pi/7) + 20*np.sin(t*2*np.pi/365)})
        with patch('app.forecast_engine._model_specs', return_value=[ModelSpec('Last observed', 'last'), ModelSpec('MSTL weekly + yearly', 'mstl')]):
            result = _fit_backtest(history, [], 14, 'daily', 'deep', 'model:MSTL weekly + yearly')
        self.assertEqual(result[4]['validation_points'], 14)
        self.assertEqual(result[4]['selection_points'], 28)
        self.assertTrue(result[4]['independent_accuracy_verified'])
        self.assertEqual(result[3]['A'], {'MSTL weekly + yearly': 1.0})

    def test_elastic_net_scaling_is_fit_only_on_training_features(self):
        X = pd.DataFrame({'price': [100., 200., 300., 400.], 'segment': ['A', 'B', 'A', 'B']})
        pipe = _make_pipeline(ElasticNet(alpha=.1, l1_ratio=.5, max_iter=10000), X)
        _fit_pipeline(pipe, X, pd.Series([10., 20., 30., 40.]))
        scale = pipe.named_steps['prep'].named_transformers_['num'].named_steps['scale']
        np.testing.assert_equal(scale.mean_, [250.])
        pipe.predict(pd.DataFrame({'price': [1000000.], 'segment': ['new']}))
        np.testing.assert_equal(scale.mean_, [250.])

    def test_nonconverged_fit_is_unavailable_not_silent_success(self):
        class Unfinished:
            def fit(self, X, y): warnings.warn('optimizer did not converge', ConvergenceWarning)
        with self.assertRaises(ConvergenceWarning): _fit_pipeline(Unfinished(), None, None)

    def test_new_methods_use_same_selection_and_confirmation_periods(self):
        history = pd.DataFrame({'item_id': 'A', 'timestamp': pd.date_range('2021-01-01', periods=60, freq='MS'), 'target': [100 + i*.5 + i%12 for i in range(60)]})
        specs = [s for s in _model_specs('fast') if s.name in {'Last observed', 'Weighted recent average', 'Holt trend', 'Holt-Winters seasonal', 'MSTL weekly + yearly'}]
        with patch('app.forecast_engine._model_specs', return_value=specs):
            result = _fit_backtest(history, [], 6, 'monthly', 'deep', 'model:Holt-Winters seasonal')
        metrics, leaderboard = result[4], result[5]
        self.assertEqual(metrics['selection_points'], 12); self.assertEqual(metrics['validation_points'], 6)
        self.assertTrue(set(metrics['selection_periods']).isdisjoint(metrics['confirmation_periods']))
        self.assertEqual(result[3]['A'], {'Holt-Winters seasonal': 1.0})
        self.assertEqual(next(r for r in leaderboard if r['model'] == 'MSTL weekly + yearly')['status'], 'unavailable')
        changed = history.copy(); changed.loc[54:, 'target'] *= 100
        with patch('app.forecast_engine._model_specs', return_value=specs):
            rerun = _fit_backtest(changed, [], 6, 'monthly', 'deep', 'model:Holt-Winters seasonal')
        self.assertEqual(result[3], rerun[3])
        for before, after in zip(leaderboard, rerun[5]):
            self.assertEqual(before['model'], after['model'])
            self.assertEqual(before['wape_pct'], after['wape_pct'])

    def test_full_fit_preserves_actual_failure_reason(self):
        history = pd.DataFrame({'item_id': 'A', 'timestamp': pd.date_range('2021-01-01', periods=36, freq='MS'), 'target': np.arange(36)+100})
        future = pd.DataFrame({'item_id': 'A', 'timestamp': pd.date_range('2024-01-01', periods=3, freq='MS')})
        spec = ModelSpec('Elastic Net + drivers', 'ml', ElasticNet())
        with patch('app.forecast_engine._fit_pipeline', side_effect=ValueError('Exact model failure')):
            predictions, _, failures = _fit_full_models_and_predict(history, future, [], 'monthly', [spec])
        self.assertFalse(predictions[spec.name])
        self.assertEqual(failures[0]['reason'], 'Exact model failure')

    def test_export_records_the_same_engine_and_method_settings(self):
        history = pd.DataFrame({'item_id': 'A', 'timestamp': pd.date_range('2021-01-01', periods=36, freq='MS'), 'target': np.arange(36)+100.})
        future = pd.DataFrame({'item_id': 'A', 'timestamp': pd.date_range('2024-01-01', periods=3, freq='MS')})
        with TemporaryDirectory() as folder, patch('app.forecast_engine._model_specs', return_value=[ModelSpec('Weighted recent average', 'weighted_average')]):
            result = run_forecast(history=history, future_covariates=future, known_driver_cols=[], horizon=3, frequency='monthly', profile='deep', runs_dir=Path(folder), method_selection='model:Weighted recent average')
            book = openpyxl.load_workbook(Path(folder)/result['run_id']/'forecast_package.xlsx', data_only=True)
            engine = dict(book['Engine'].iter_rows(min_row=2, values_only=True))
            settings = {name: json.loads(values) for name, values in book['Method settings'].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(engine, result['engine']); self.assertEqual(settings, result['method_settings'])
            self.assertEqual(settings['weighted_average']['weights_oldest_to_newest'], [1, 2, 3])
            book.close()

    def test_family_ensemble_must_beat_its_best_eligible_member(self):
        history = pd.DataFrame({'item_id': 'A', 'timestamp': pd.date_range('2021-01-01', periods=60, freq='MS'), 'target': [100.] * 60})
        specs = [ModelSpec('Seasonal naive', 'naive'), ModelSpec('AutoETS', 'ets'), ModelSpec('Last observed', 'last')]
        def forecast(spec, history, steps, lag, frequency=None):
            return [100. if spec.kind in {'naive', 'last'} else 105.] * steps
        with patch('app.forecast_engine._model_specs', return_value=specs), patch('app.forecast_engine._stat_forecast', side_effect=forecast):
            result = _fit_backtest(history, [], 6, 'monthly', 'deep', 'seasonal')
        self.assertEqual(result[3]['A'], {'Seasonal naive': 1.0})
        self.assertEqual(result[4]['ensemble_gate_series'], 1)


if __name__ == '__main__': unittest.main()
